from __future__ import annotations

import asyncio
from copy import deepcopy
import io
import json
import threading

from docx import Document
from fastapi import HTTPException
from fastapi.testclient import TestClient
from pypdf import PdfReader
import httpx
import pytest

from audio_intel import api, db, optimization_api as routes, optimization_engine as engine, optimization_store as store, optimization_summaries as insights
from audio_intel.transcript_export import TranscriptExportRequest, optimized_export_result, format_transcript
from audio_intel import transcript_docx, transcript_pdf
from test_optimization import configured, local, job, source, start, url, wait, PROFILE


def test_summary_round_precedes_publication_and_uses_source(configured, monkeypatch):
    entered, release = threading.Event(), threading.Event()
    calls = []
    raw = source(2)
    raw['video'] = {'scenes': [{'id': 'a', 'start': 0, 'end': 6, 'image': None}]}
    async def generate(profile, text, kind):
        calls.append((profile.copy(), text, kind))
        entered.set()
        while not release.is_set():
            await asyncio.sleep(.01)
        return ['预算未确认，周五测试。', '阿特拉斯预算15万元。'] if kind == 'group' else ['预算测试安排']
    monkeypatch.setattr(insights, 'generate', generate)
    identifier = job(raw)
    with TestClient(api.create_app()) as client:
        assert start(client, identifier).status_code == 202
        assert entered.wait(3)
        state = client.get(url(identifier)).json()
        run = state['candidate']['id']
        assert state['published'] is None
        assert state['candidate']['phase'] == 'summaries'
        assert client.get(url(identifier)+'/result', params={'run_id': run}).status_code == 409
        release.set()
        wait(client, identifier, lambda s: bool(s['published']))
        result = client.get(url(identifier)+'/result', params={'run_id': run}).json()
        assert result['summary_status']['state'] == 'completed'
        assert len(result['paragraphs'][0]['group_summary']['sentences']) == 2
        assert result['scenes'][0]['title_summary']['title'] == '预算测试安排'
        light = client.get(url(identifier)+'/summaries', params={'run_id': run})
        assert client.get(url(identifier)+'/summaries', params={'run_id': run}, headers={'If-None-Match': light.headers['etag']}).status_code == 304
        assert 'paragraphs' not in light.json()
        metadata = client.get('/api/v1/jobs').text
        assert '预算测试安排' not in metadata
        assert '预算未确认，周五测试。' not in metadata
    assert len(calls) == 3
    assert calls[-1][1] == '\n'.join(u['original_text'] for u in store.units(run))
    assert all(c[0]['context_length'] == 32768 and c[0]['model_digest'] == PROFILE['model_digest'] for c in calls)
    assert '那个' in calls[0][1]  # Inputs are immutable sources, not rewritten prose.


@pytest.mark.parametrize('failure, expected_calls', [('transport', 3), ('context', 1), ('invalid', 5)])
def test_auxiliary_failures_publish_prose_and_bound_work(configured, monkeypatch, failure, expected_calls):
    calls = []
    async def fail(*args):
        calls.append(args)
        if failure == 'transport':
            raise httpx.ConnectError('down')
        if failure == 'context':
            raise engine.ContextWindowError('allocated 16K below requested 64K')
        raise engine.InvalidOutput('unverified facts')
    monkeypatch.setattr(insights, 'generate', fail)
    identifier = job(source(5))
    with TestClient(api.create_app()) as client:
        start(client, identifier)
        run = wait(client, identifier, lambda s: bool(s['published']))['published']
        assert run['state'] == 'succeeded'
        assert run['summary_status']['state'] == 'partial'
        assert run['summary_status']['counts']['group']['failed'] == 5
        assert all(u['text'] for u in store.units(run['id']))
    assert len(calls) == expected_calls


def test_retry_preserves_successes_prose_idempotency_and_snapshots(configured, monkeypatch):
    calls = []
    async def generate(profile, text, kind):
        calls.append((profile.copy(), text))
        if len(calls) == 2:
            raise engine.InvalidOutput('bad')
        return [] if '第2项' in text else ['预算尚未确认。']
    monkeypatch.setattr(insights, 'generate', generate)
    identifier = job(source(3))
    with TestClient(api.create_app()) as client:
        start(client, identifier)
        run = wait(client, identifier, lambda s: bool(s['published']))['published']['id']
        before = store.units(run)
        assert insights.read(run)['groups']['2']['state'] == 'not_applicable'
        store.save_settings({**PROFILE, 'model': 'different-global-model', 'context_length': 65536})
        root = url(identifier)+'/summaries/retry'
        headers = {'Idempotency-Key': 'summary-retry-key-1'}
        assert client.post(root, json={'run_id': run}).status_code == 422
        assert client.post(root, json={'run_id': run}, headers=headers).status_code == 202
        wait(client, identifier, lambda s: s['published']['summary_status']['state'] == 'completed')
        assert len(calls) == 4
        assert calls[-1][0]['model'] == PROFILE['model']
        assert calls[-1][0]['context_length'] == 32768
        assert store.units(run) == before
        store.save_settings({**PROFILE, 'enabled': False})
        assert client.post(root, json={'run_id': run}, headers=headers).status_code == 200
        store.save_settings(PROFILE)
        assert client.post(root, json={'run_id': run}, headers={'Idempotency-Key': 'summary-no-work-key'}).status_code == 200
        assert len(calls) == 4
        other = job(source(1))
        assert client.get(url(other)+'/summaries', params={'run_id': run}).status_code == 404
        with db.connect() as c:
            assert c.execute('SELECT count(*) FROM optimization_summary_retries').fetchone()[0] == 2
        assert client.delete(url(identifier)+'?purge=true').status_code == 204
        with db.connect() as c:
            for table in ('optimization_summaries', 'optimization_summary_fragments', 'optimization_summary_retries'):
                assert not c.execute(f'SELECT 1 FROM {table}').fetchall()


def test_pause_retry_rejects_late_response_and_restart_retains_items(configured, monkeypatch):
    async def bad(*args):
        raise engine.InvalidOutput('bad')
    monkeypatch.setattr(insights, 'generate', bad)
    identifier = job(source(2))
    entered = threading.Event()
    async def late(*args):
        entered.set()
        try:
            await asyncio.sleep(30)
        except asyncio.CancelledError:
            return ['不应写入的旧响应']
    with TestClient(api.create_app()) as client:
        start(client, identifier)
        run = wait(client, identifier, lambda s: bool(s['published']))['published']['id']
        monkeypatch.setattr(insights, 'generate', late)
        client.post(url(identifier)+'/summaries/retry', json={'run_id': run}, headers={'Idempotency-Key': 'pause-summary-retry'})
        assert entered.wait(3)
        assert client.post(url(identifier)+'/summaries/pause', json={'run_id': run}).status_code == 200
        assert insights.read(run)['summary_status']['state'] == 'paused'
        assert all(not v['sentences'] for v in insights.read(run)['groups'].values())
    with TestClient(api.create_app()) as client:
        assert client.get(url(identifier)).json()['published']['summary_status']['state'] == 'paused'
        assert store.get(run)['published']


def test_v19_migration_is_additive_and_legacy_runs_do_not_generate(configured):
    identifier = job(source(1))
    run, _ = store.create(identifier, 'legacy-key', False, PROFILE, engine.source_hash(source(1)), ['Atlas'], engine.prepare_units(source(1)))
    store.transition(run, 'running', ('queued',))
    asyncio.run(routes.Manager(configured).execute(run))
    before = store.units(run)
    with db.connect() as c:
        c.executescript('DROP TABLE optimization_summary_retries; DROP TABLE optimization_summary_fragments; DROP TABLE optimization_summaries; DROP INDEX idx_optimization_summary_queue;')
        for column in ('summary_version', 'summary_state', 'summary_counts_json', 'summary_revision', 'summary_requested_at'):
            c.execute(f'ALTER TABLE transcript_optimizations DROP COLUMN {column}')
        c.execute('UPDATE schema_meta SET version=19')
    db.init_db();db.init_db()
    assert store.units(run) == before
    assert store.get(run)['published']
    assert insights.read(run)['summary_status']['state'] == 'not_requested'
    with TestClient(api.create_app()) as client:
        assert client.post(url(identifier)+'/summaries/retry', json={'run_id': run}, headers={'Idempotency-Key': 'legacy-summary-retry'}).status_code == 409
    with db.connect() as c:
        assert c.execute('SELECT version FROM schema_meta').fetchone()[0] == 21
        assert not c.execute('PRAGMA foreign_key_check').fetchall()


@pytest.mark.parametrize('text,kind,valid', [
    ('中'*50,'group',True), ('中'*51,'group',False), ('日'*20,'scene',True), ('日'*21,'scene',False),
    (' '.join(['word']*30),'group',True), (' '.join(['word']*31),'group',False),
    (' '.join(['word']*10),'scene',True), (' '.join(['word']*11),'scene',False),
    ('中'*25+' '+' '.join(['word']*15),'group',True), ('中'*26+' '+' '.join(['word']*15),'group',False),
    ('预算15万元，周五测试。','group',True), ('a\nb','group',False), ('','group',False),
])
def test_multilingual_length_budget(text, kind, valid):
    assert insights.fits_length(text, kind) is valid


@pytest.mark.parametrize('candidate', [
    {'texts':['预算50万元。'],'evidence':['预算15万元。'],'no_content':False},
    {'texts':['中'*51],'evidence':['预算15万元。'],'no_content':False},
    {'texts':['预算15万元。'],'evidence':['不是来源'],'no_content':False},
    {'texts':['预算15万元。','预算15万元。'],'evidence':['预算15万元。'],'no_content':False},
    {'texts':[42],'evidence':[],'no_content':False},
])
def test_generation_rejects_invalid_content_in_two_rounds(monkeypatch, candidate):
    calls=[]
    async def structured(profile, payload):
        calls.append(payload)
        return candidate
    monkeypatch.setattr(engine, 'structured', structured)
    with pytest.raises(engine.InvalidOutput):
        asyncio.run(insights.generate(PROFILE, '预算15万元。', 'group'))
    assert len(calls) == 2
    assert calls[0]['truncate'] is False and calls[0]['think'] is False


def test_fact_priority_check_and_no_content_are_independent(monkeypatch):
    calls=[]
    async def structured(profile, payload):
        calls.append(payload)
        if payload['format'] == insights.CHECK_SCHEMA:
            return {'valid': len(calls) == 4}
        return {'texts': [],'evidence': [],'no_content': True}
    monkeypatch.setattr(engine, 'structured', structured)
    assert asyncio.run(insights.generate(PROFILE, '好的，谢谢。', 'group')) == []
    assert len(calls) == 4


def test_scene_reduction_covers_all_source_and_resumes_checkpoints(configured, monkeypatch):
    identifier = job(source(1))
    profile = {**PROFILE, 'summary_version': 1}
    run, _ = store.create(identifier,'scene-cache-key',False,profile,engine.source_hash(source(1)),[],engine.prepare_units(source(1)))
    with db.connect() as c:
        c.execute("UPDATE transcript_optimizations SET summary_state='running' WHERE id=?", (run,))
    monkeypatch.setattr(engine, 'request_fits', lambda p,r: len(json.loads(r['messages'][1]['content'])['source']) <= 200)
    raw='甲'*200+'乙'*200+'丙'*200+'丁'*200
    captured=[]
    async def generate(profile,text,kind):
        if kind=='fragment':
            captured.append(text)
            if len(captured)==2:
                raise asyncio.CancelledError()
            return [''.join(dict.fromkeys(text))+'主题']
        assert all(char in text for char in '甲乙丙丁')
        return ['完整画面主题']
    monkeypatch.setattr(insights, 'generate', generate)
    with pytest.raises(asyncio.CancelledError):
        asyncio.run(insights.scene_title(run,profile,raw))
    captured.clear()
    async def resume(profile,text,kind):
        if kind=='fragment':
            captured.append(text)
            return [''.join(dict.fromkeys(text))+'主题']
        assert all(char in text for char in '甲乙丙丁')
        return ['完整画面主题']
    monkeypatch.setattr(insights, 'generate', resume)
    assert asyncio.run(insights.scene_title(run,profile,raw)) == ['完整画面主题']
    assert ''.join(captured) == raw[200:]


def test_exports_include_redacted_summaries_once_below_comparison(configured, monkeypatch):
    async def generate(profile,text,kind):
        return ['阿特拉斯预算15万元。'] if kind=='group' else ['阿特拉斯预算']
    monkeypatch.setattr(insights,'generate',generate)
    raw=source(1);raw['video']={'scenes':[{'id':'a','start':0,'end':3,'image':None}]}
    identifier=job(raw)
    with TestClient(api.create_app()) as client:
        start(client,identifier)
        run=wait(client,identifier,lambda s:bool(s['published']))['published']['id']
    policy=TranscriptExportRequest(format='docx',text_version='comparison',optimization_id=run,replacements=[{'source':'阿特拉斯','target':'项目X'}])
    result=optimized_export_result(db.get_job(identifier),policy)
    assert result['segments'][0]['group_summary']['sentences']==['项目X预算15万元。']
    assert result['video']['scenes'][0]['title_summary']['title']=='项目X预算'
    assert '阿特拉斯' not in json.dumps(result,ensure_ascii=False)
    payload=transcript_docx.render('Meeting',result,configured.data_dir)
    document=Document(io.BytesIO(payload))
    text='\n'.join(p.text for p in document.paragraphs)
    assert text.count('小结 / Summary：项目X预算15万元。')==1
    assert '项目X预算' in text
    assert all('小结' not in cell.text for table in document.tables for row in table.rows for cell in row.cells)
    pdf=PdfReader(io.BytesIO(transcript_pdf.render('Meeting',result,configured.data_dir)))
    text='\n'.join(page.extract_text() for page in pdf.pages)
    assert '项目X预算15万元' in text and '小结' in text
    assert 'Summary' not in format_transcript(result,'txt')


def test_retry_admission_and_light_reads(configured, monkeypatch):
    identifier=job(source(1))
    profile={**PROFILE,'summary_version':1}
    run,_=store.create(identifier,'admission-source-key',False,profile,engine.source_hash(source(1)),['Atlas'],engine.prepare_units(source(1)))
    store.transition(run,'running',('queued',))
    asyncio.run(routes.Manager(configured).execute(run))
    with db.connect() as c:
        c.execute("UPDATE optimization_summaries SET state='unavailable',value_json='[]' WHERE run_id=?",(run,))
        insights.update_counts(c,run,'partial')
    for n in range(5):
        other=job(source(1))
        store.create(other,f'queued-key-{n}',False,PROFILE,engine.source_hash(source(1)),[],engine.prepare_units(source(1)))
    with pytest.raises(HTTPException) as error:
        insights.retry(run,'full-queue-key',0)
    assert error.value.status_code==429 and error.value.headers['Retry-After']=='5'
    with db.connect() as c:
        c.execute("DELETE FROM transcript_optimizations WHERE state='queued'")
    monkeypatch.setattr(insights.shutil,'disk_usage',lambda p:type('Space',(),{'free':0})())
    with pytest.raises(HTTPException) as error:
        insights.retry(run,'disk-full-key',1)
    assert error.value.status_code==429
    with db.connect() as c:
        assert not c.execute('SELECT 1 FROM optimization_summary_retries').fetchall()
    with TestClient(api.create_app()) as client:
        def no_result_read(*args,**kwargs):
            raise AssertionError('Light summary reads must not load result_json')
        monkeypatch.setattr(db,'get_job',no_result_read)
        assert client.get(url(identifier)+'/summaries',params={'run_id':run}).status_code==200


def test_rejected_spoken_number_does_not_become_approved_previous_budget(monkeypatch):
    calls=[]
    async def structured(profile,payload):
        calls.append(payload)
        return {'texts':['预算从50万元调整为15万元。'],'evidence':['刚才说50万元不对，实际15万元。'],'no_content':False}
    monkeypatch.setattr(engine,'structured',structured)
    with pytest.raises(engine.InvalidOutput,match='corrected or rejected'):
        asyncio.run(insights.generate(PROFILE,'刚才说50万元不对，实际15万元。','group'))
    assert len(calls)==2


def test_scene_change_cascades_insights_and_retry_receipts(configured):
    raw=source(1);raw['video']={'scenes':[{'id':'a','start':0,'end':3,'image':None}]}
    identifier=job(raw)
    profile={**PROFILE,'summary_version':1,'summary_scene_ids':['a']}
    run,_=store.create(identifier,'scene-invalidation-key',False,profile,engine.source_hash(raw),['Atlas'],engine.prepare_units(raw))
    store.transition(run,'running',('queued',))
    asyncio.run(routes.Manager(configured).execute(run))
    with db.connect() as c:
        c.execute('INSERT INTO optimization_summary_retries VALUES(?,?,?)',(identifier,'key',run))
        c.execute('INSERT INTO optimization_summary_fragments VALUES(?,?,?)',(run,'hash','[]'))
        after=deepcopy(raw);after['video']['scenes'][0]['end']=2.5
        c.execute('BEGIN IMMEDIATE')
        assert store.invalidate_scene_change(c,identifier,raw,after)
        c.commit()
        for table in ('optimization_summaries','optimization_summary_fragments','optimization_summary_retries'):
            assert not c.execute(f'SELECT 1 FROM {table}').fetchall()
        assert not c.execute('PRAGMA foreign_key_check').fetchall()
