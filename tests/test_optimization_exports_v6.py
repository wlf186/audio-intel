"""Published content remains exportable while independent review progresses."""
from copy import deepcopy
import io

from docx import Document
from pypdf import PdfReader
import pytest

from audio_intel import transcript_docx, transcript_pdf
from audio_intel.transcript import reading_sections
from audio_intel.transcript_export import TranscriptExportRequest, export_result, format_transcript, optimized_export_result


def draft():
    return {'prompt_version': 6, 'complete': True, 'review_state': 'running', 'uncertain_units': 0,
            'unavailable_units': 0, 'basic_cleanup_units': 0,
            'paragraphs': [{'id': 'g1', 'ordinal': 0, 'speaker': 's1', 'speaker_label': 'Mike',
                'start': 0, 'end': 20, 'text': '请发布Atlas模型。', 'original_text': '那个发布阿特拉斯模型。',
                'omitted': False, 'candidate_text': None, 'review_status': None,
                'review_outcome': 'pending', 'generation_method': 'model', 'merged_source_count': 2,
                'scene_ids': ['a', 'b'], 'segment_ids': [1, 2],
                'sources': [{'segment_id': 1, 'start': 0, 'end': 2, 'text': '那个发布阿特拉斯'},
                            {'segment_id': 2, 'start': 2, 'end': 20, 'text': '模型。'}]}]}


def job():
    return {'id': 'job', 'result': {'speakers': [{'id': 's1', 'name': 'Mike', 'label': 'Mike'}],
        'segments': [], 'video': {'scenes': [
            {'id': 'a', 'start': 0, 'end': 2, 'image': None},
            {'id': 'b', 'start': 2, 'end': 20, 'image': None}]}}}


def exported(monkeypatch, version='comparison', value=None):
    from audio_intel import optimization_engine, optimization_store
    data = value or draft()
    monkeypatch.setattr(optimization_store, 'get', lambda _: {'id': 'a' * 32, 'job_id': 'job',
        'published': True, 'state': 'succeeded', 'config': {'prompt_version': 6}})
    monkeypatch.setattr(optimization_engine, 'reading_result', lambda *args, **kwargs: deepcopy(data))
    return optimized_export_result(job(), TranscriptExportRequest(format='pdf', text_version=version,
        optimization_id='a' * 32, replacements=[{'source': '阿特拉斯', 'target': '项目甲'}, {'source': 'Atlas', 'target': '项目甲'}]))


@pytest.mark.parametrize('state', ['queued', 'running', 'paused', 'unavailable'])
def test_published_draft_exports_before_review_finishes(monkeypatch, state):
    data = draft(); data['review_state'] = state
    result = exported(monkeypatch, value=data)
    assert 'not complete' in result['review_notice']
    assert result['segments'][0]['text'] == '请发布项目甲模型。'
    assert result['segments'][0]['sources'][0]['text'] == '那个发布项目甲'
    assert '阿特拉斯' not in str(result) and 'Atlas' not in str(result)
    sections = reading_sections(result)
    assert len(sections[0]['paragraphs']) == 1 and not sections[1]['paragraphs']
    assert sections[1]['linked_paragraphs'] == [{'start': 0, 'scene_number': 1}]


@pytest.mark.parametrize('format', ['pdf', 'docx', 'txt'])
def test_review_and_cleanup_labels_survive_redaction(monkeypatch, tmp_path, format):
    data = draft(); data.update(review_state='completed', basic_cleanup_units=1)
    data['paragraphs'][0].update(review_status='needs_review', review_outcome='fact',
        generation_method='basic_cleanup', review_details={'stage': 'review', 'issues': [
            {'source_quote': '阿特拉斯', 'candidate_quote': 'Atlas', 'reason': '请核对阿特拉斯名称'}]})
    result = exported(monkeypatch, version='optimized' if format == 'txt' else 'comparison', value=data)
    if format == 'pdf':
        text = '\n'.join(p.extract_text() for p in PdfReader(io.BytesIO(transcript_pdf.render('title', result, tmp_path))).pages)
    elif format == 'docx':
        document = Document(io.BytesIO(transcript_docx.render('title', result, tmp_path)))
        text = '\n'.join(p.text for p in document.paragraphs) + '\n'.join(c.text for t in document.tables for r in t.rows for c in r.cells)
        # v6 is reading-first and must not strike out whole original paragraphs.
        assert not any(run.font.strike for t in document.tables for row in t.rows for cell in row.cells for p in cell.paragraphs for run in p.runs)
    else:
        text = format_transcript(result, 'txt')
    assert 'Needs review' in text and ('Basic cleanup' in text or 'basic cleanup' in text)
    assert '项目甲' in text and 'Atlas' not in text and '阿特拉斯' not in text
    if format != 'txt':
        assert 'Merged sources' in text and '00:00:02' in text
        assert '请核对项目甲名称' in text
        assert text.count('请发布项目甲模型。') == 1


def test_completed_review_snapshot_changes_export_cache_key(monkeypatch):
    before = exported(monkeypatch)
    data = draft(); data['review_state'] = 'completed'; data['paragraphs'][0]['review_outcome'] = 'pass'
    after = exported(monkeypatch, value=data)
    assert before['review_notice'] and not after['review_notice']
    assert before != after
