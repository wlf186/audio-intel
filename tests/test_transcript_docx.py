from copy import deepcopy
from dataclasses import replace
import io
import zipfile

from docx import Document
from fastapi.testclient import TestClient
from PIL import Image
import pytest

from audio_intel import api, db, transcript_docx as docx
from test_voiceprint_groups_exports import local, transcript


def text(payload):
    document = Document(io.BytesIO(payload))
    return '\n'.join(p.text for p in document.paragraphs) + '\n'.join(cell.text for table in document.tables for row in table.rows for cell in row.cells)


def test_audio_merges_same_id_preserves_words_and_special_characters(tmp_path):
    result = transcript()
    result['segments'].extend([
        {'id':1,'speaker':'Speaker_0','speaker_label':'Vince','start':3,'end':4,'text':'嗯，<xml> & "中文"\n下一行。'},
        {'id':2,'speaker':'Speaker_1','speaker_label':'Vince','start':4,'end':5,'text':'另一个人。'},
    ])
    before = deepcopy(result)
    payload = docx.render('测试', result, tmp_path)
    document = Document(io.BytesIO(payload))
    labels = [p for p in document.paragraphs if ' – ' in p.text]
    assert len(labels) == 2
    assert '00:00:04' in labels[0].text
    assert '嗯，<xml> & "中文"\n下一行。' in text(payload)
    assert result == before


def test_video_docx_visibility_cache_redaction_and_images(local):
    result = transcript()
    result.pop('hotword_context'); result.pop('voiceprint_library')
    for speaker in result['speakers']:
        speaker['label'] = speaker['name']
        speaker.pop('voiceprint_match', None)
    result['video'] = {'contract_version':1, 'time_origin':0, 'detector':'presentation-v1', 'revision':1, 'duration':6, 'scenes':[
        {'id':'first','start':0,'end':3,'frame_time':1,'image':'private-red.jpg'},
        {'id':'second','start':3,'end':6,'frame_time':4,'image':'private-blue.jpg'},
    ]}
    result['video']['original_scenes'] = deepcopy(result['video']['scenes'])
    job = db.create_job('asr', '私密标题.mp4', {})
    folder = local.jobs_dir / job['id'] / 'output'
    frames = folder / 'frames'; frames.mkdir(parents=True)
    for name, color in [('private-red.jpg','red'),('private-blue.jpg','blue')]:
        Image.new('RGB',(320,180),color).save(frames/name)
    raw = {p.name:p.read_bytes() for p in frames.iterdir()}
    db.update_job(job['id'],state='succeeded',result_json=result)
    root = f"/api/v1/jobs/{job['id']}"
    with TestClient(api.create_app()) as client:
        assert 'docx' in client.get('/api/v1/capabilities').json()['asr']['transcript_export']['formats']
        original = client.post(root+'/transcript-export',json={'format':'docx'}).content
        assert len(Document(io.BytesIO(original)).inline_shapes) == 2
        saved = client.patch(root+'/video-scenes',json={'action':'set_export_visibility','revision':1,'export_hidden_scene_ids':['first']})
        assert saved.status_code == 200
        for redact in (False,True):
            response = client.post(root+'/transcript-export',json={'format':'docx','redact':redact,'replacements':[{'source':'广东移动','target':'替换内容'}]})
            assert response.status_code == 200
            assert response.headers['content-type'] == 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'
            assert '.docx' in response.headers['content-disposition']
            assert len(response.content) == int(response.headers['content-length'])
            document = Document(io.BytesIO(response.content))
            assert len(document.inline_shapes) == 1
            content = text(response.content)
            assert 'Picture hidden' in content and '此段无识别语音' in content
            assert ('替换内容' if redact else '广东移动') in content
            with zipfile.ZipFile(io.BytesIO(response.content)) as package:
                media = [package.read(name) for name in package.namelist() if name.startswith('word/media/')]
                assert raw['private-red.jpg'] not in media
                assert raw['private-blue.jpg'] in media
                xml = b''.join(package.read(name) for name in package.namelist() if name.endswith('.xml'))
                assert b'private-red' not in xml and b'private-blue' not in xml
                if redact:
                    for secret in ('私密标题','Vince Carter','01234567','广东移动'):
                        assert secret.encode() not in xml
        assert client.patch(root+'/video-scenes',json={'action':'set_export_visibility','revision':2,'export_hidden_scene_ids':[]}).status_code == 200
        restored = client.post(root+'/transcript-export',json={'format':'docx'}).content
        assert len(Document(io.BytesIO(restored)).inline_shapes) == 2
        assert 'Picture hidden' not in text(restored)
    assert {p.name:p.read_bytes() for p in frames.iterdir()} == raw
    assert db.get_job(job['id'])['result']['segments'] == result['segments']


def test_docx_budgets_paths_and_busy_response(local, monkeypatch):
    job = db.create_job('asr','budget.wav',{})
    result = transcript(); db.update_job(job['id'],state='succeeded',result_json=result)
    root = f"/api/v1/jobs/{job['id']}/transcript-export"
    docx._cache.clear()
    with TestClient(api.create_app()) as client:
        with monkeypatch.context() as patch:
            patch.setattr(docx,'_max_text_chars',1)
            assert client.post(root,json={'format':'docx'}).status_code == 413
        with docx._render_lock:
            response = client.post(root,json={'format':'docx'})
            assert response.status_code == 429 and response.headers['retry-after'] == '2'
        assert client.post(root,json={'format':'docx'}).status_code == 200
    result['video'] = {'scenes':[{'id':'s','start':0,'end':3,'frame_time':1,'image':'../outside.jpg'}]}
    with pytest.raises(ValueError,match='Invalid frame'):
        docx.render('path',result,local.jobs_dir)
    frames=local.jobs_dir/'frames';frames.mkdir()
    Image.new('RGB',(20,20),'red').save(frames/'one.jpg')
    result['video']['scenes'][0]['image']='one.jpg'
    monkeypatch.setattr(docx,'_max_image_bytes',1)
    with pytest.raises(ValueError,match='Images exceed'):
        docx.render('budget',result,local.jobs_dir)
    result['video']['export_hidden_scene_ids']=['s']
    (frames/'one.jpg').unlink()
    assert 'Picture hidden' in text(docx.render('hidden',result,local.jobs_dir))


def test_docx_authentication(local, monkeypatch):
    monkeypatch.setattr(api,'settings',replace(local,api_key='docx-test'))
    job=db.create_job('asr','auth.wav',{});db.update_job(job['id'],state='succeeded',result_json=transcript())
    root=f"/api/v1/jobs/{job['id']}/transcript-export"
    with TestClient(api.create_app()) as client:
        assert client.post(root,json={'format':'docx'}).status_code==401
        assert client.post(root,json={'format':'docx'},headers={'Authorization':'Bearer docx-test'}).status_code==200
        assert client.post('/api/v1/auth/session',headers={'Authorization':'Bearer docx-test'}).status_code==204
        assert client.post(root,json={'format':'docx'}).status_code==403
        assert client.post(root,json={'format':'docx'},headers={'Origin':'http://testserver'}).status_code==200
