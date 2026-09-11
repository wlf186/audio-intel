from __future__ import annotations

import datetime as dt
import io
import time
import uuid
import zipfile
from pathlib import Path

import pytest

from audio_intel.document_text import extract, segment, digest, _plain_headings
from audio_intel.document_office import display_cell, validate_package
from test_documents import document_client, submit_document


def rewrite_member(path, member, transform):
    with zipfile.ZipFile(path) as archive:
        entries = [(i, archive.read(i)) for i in archive.infolist()]
    with zipfile.ZipFile(path, 'w') as archive:
        for info, data in entries:
            archive.writestr(info, transform(data) if info.filename == member else data)


@pytest.fixture
def office_files(tmp_path):
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation
    from pptx.util import Inches
    files = []
    doc = Document()
    doc.add_heading('第一章', 1)
    doc.add_paragraph('这是第一章的正文。')
    doc.add_heading('第二章', 1)
    doc.add_paragraph('这是第二章的正文。')
    path = tmp_path / 'example.docx'
    doc.save(path)
    files.append(path)
    book = Workbook()
    book.active.title = '第一张表'
    book.active.append(['资产', '黄金'])
    book.create_sheet('第二张表').append(['状态', '正常'])
    path = tmp_path / 'example.xlsx'
    book.save(path)
    files.append(path)
    slides = Presentation()
    for i in (1, 2):
        slide = slides.slides.add_slide(slides.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(1)).text = f'第{i}页正文。'
    path = tmp_path / 'example.pptx'
    slides.save(path)
    files.append(path)
    return files


def test_word_custom_outline_fields_tables_and_final_text(tmp_path):
    from docx import Document
    from docx.enum.style import WD_STYLE_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document()
    base = doc.styles.add_style('自定义主标题', WD_STYLE_TYPE.PARAGRAPH)
    outline = OxmlElement('w:outlineLvl'); outline.set(qn('w:val'), '0')
    base.element.get_or_add_pPr().append(outline)
    derived = doc.styles.add_style('继承的标题', WD_STYLE_TYPE.PARAGRAPH)
    derived.base_style = base
    p = doc.add_paragraph()
    for kind in ('begin', 'separate'):
        r = p.add_run()
        if kind == 'begin':
            field = OxmlElement('w:fldChar'); field.set(qn('w:fldCharType'), kind); r._r.append(field)
            ins = OxmlElement('w:instrText'); ins.text = ' TOC \\o "1-3" '; r._r.append(ins)
        else:
            field = OxmlElement('w:fldChar'); field.set(qn('w:fldCharType'), kind); r._r.append(field)
    doc.add_paragraph('DUPLICATE_NAVIGATION')
    field = OxmlElement('w:fldChar'); field.set(qn('w:fldCharType'), 'end')
    doc.add_paragraph().add_run()._r.append(field)
    doc.add_paragraph('主章一', derived)
    doc.add_paragraph('before table')
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = 'merged once'
    table.cell(0, 0).merge(table.cell(1, 0))
    table.cell(0, 1).text = 'right one'
    table.cell(1, 1).text = 'right two'
    doc.add_paragraph('after table')
    p = doc.add_paragraph('visible ')
    deleted = OxmlElement('w:del')
    r = OxmlElement('w:r'); text = OxmlElement('w:t'); text.text = 'DELETED'; r.append(text); deleted.append(r); p._p.append(deleted)
    inserted = OxmlElement('w:ins')
    r = OxmlElement('w:r'); text = OxmlElement('w:t'); text.text = 'INSERTED'; r.append(text); inserted.append(r); p._p.append(inserted)
    doc.add_paragraph('主章二', derived)
    doc.add_paragraph('正文。' * 18000)
    source = tmp_path / 'outline.docx'; doc.save(source)
    parsed = extract(source, 5000000, 256*1024**2)
    assert 'DUPLICATE_NAVIGATION' not in parsed['text']
    assert parsed['text'].count('merged once') == 1
    assert parsed['text'].index('before table') < parsed['text'].index('merged once') < parsed['text'].index('after table')
    assert 'DELETED' not in parsed['text'] and 'INSERTED' in parsed['text']
    assert [s['title'] for s in segment(parsed)['sections']] == ['主章一', '主章二']
    assert segment(parsed)['sections'][1]['char_count'] > 50000
    split = segment(parsed, 'length', 1000)
    assert ''.join(parsed['text'][s['start']:s['end']] for s in split['sections']) == parsed['text']


def test_excel_headers_formulas_formats_hidden_and_sparse(tmp_path):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.worksheet.table import Table
    book = Workbook()
    sheet = book.active; sheet.title = '数据'
    sheet.append(['封面指标', '第一条值', '说明'])
    sheet.append([])
    sheet.append(['资产', '涨跌', '日期', '编号', '缺失'])
    for cell in sheet[3]: cell.font = Font(bold=True)
    sheet.append(['黄金', '=1/4', dt.datetime(2026, 9, 11), 12, '=1+1'])
    sheet['B4'].number_format = '0.00%'
    sheet['C4'].number_format = 'yyyy-mm-dd'
    sheet['D4'].number_format = '000000'
    sheet.append(['HIDDEN_ROW', 0]); sheet.row_dimensions[5].hidden = True
    sheet['F4'] = 'HIDDEN_COLUMN'; sheet.column_dimensions['F'].hidden = True
    sheet['A1048576'] = '最后一行'
    other = book.create_sheet('正式表格')
    other.append(['名称', '结果']); other.append(['测试', False])
    other.add_table(Table(displayName='Declared', ref='A1:B2'))
    hidden = book.create_sheet('隐藏表'); hidden['A1'] = 'HIDDEN_SHEET'; hidden.sheet_state = 'hidden'
    path = tmp_path/'data.xlsx'; book.save(path)
    rewrite_member(path, 'xl/worksheets/sheet1.xml', lambda b: b.replace(b'<f>1/4</f><v></v>', b'<f>1/4</f><v>0.25</v>'))
    start = time.monotonic()
    parsed = extract(path, 5000000, 256*1024**2)
    assert time.monotonic() - start < 5
    text = parsed['text']
    assert '封面指标：第一条值；说明' in text
    assert '资产：黄金；涨跌：25.00%；日期：2026-09-11；编号：000012；缺失：公式结果缺失' in text
    assert '结果：FALSE' in text and '最后一行' in text
    assert 'HIDDEN_' not in text
    assert any('数据!E4' in w and 'recalculate' in w for w in parsed['warnings'])
    assert [s['title'] for s in segment(parsed)['sections']] == ['数据', '正式表格']


@pytest.mark.parametrize('value,fmt,expected', [
    (-.002348848618, '0.00%', '-0.23%'), (6.71, '#,##0.0000', '6.7100'),
    (1234.5, '#,##0.00', '1,234.50'), (1.2, '0.0#', '1.2'),
    (-2.5, '0.00;(0.00)', '(2.50)'), (123000, '0,', '123'),
    (0, '0.00', '0.00'), (False, 'General', 'FALSE'),
])
def test_excel_display(value, fmt, expected):
    from openpyxl import Workbook
    c = Workbook().active['A1']; c.value = value; c.number_format = fmt
    warnings = []
    assert display_cell(c, warnings, 'A1') == expected
    assert not warnings


def test_excel_merged_header_and_unsupported_format(tmp_path):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    book = Workbook(); s = book.active
    s.append(['分类', None, '值']); s.merge_cells('A1:B1')
    s['A1'].font = s['C1'].font = Font(bold=True)
    s.append(['甲', '乙', 1.5]); s['C2'].number_format = '0.00E+00'
    path = tmp_path/'merged.xlsx'; book.save(path)
    parsed = extract(path, 5000000, 256*1024**2)
    assert '分类：甲；分类：乙；值：1.5' in parsed['text']
    assert any('unsupported number format' in w for w in parsed['warnings'])


def test_slides_visual_order_groups_notes_hidden_and_tables(tmp_path):
    from pptx import Presentation
    from pptx.util import Inches
    slides = Presentation(); slide = slides.slides.add_slide(slides.slide_layouts[6])
    # Deliberately reverse z-order; reading must follow geometry.
    for y, x, text in [(2, 5, '④'), (2, 1, '③'), (1, 5, '②'), (1, 1, '①')]:
        slide.shapes.add_textbox(Inches(x), Inches(y), Inches(2), Inches(.4)).text = text
    group = slide.shapes.add_group_shape()
    group.shapes.add_textbox(Inches(1), Inches(3), Inches(2), Inches(.4)).text = 'GROUP'
    table = slide.shapes.add_table(2, 2, Inches(1), Inches(4), Inches(6), Inches(1)).table
    table.cell(0, 0).text = 'MERGED'; table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(1, 0).text = 'BOTTOM'
    slide.notes_slide.notes_text_frame.text = 'NOTES_OMITTED'
    slides.slides.add_slide(slides.slide_layouts[6])
    hidden = slides.slides.add_slide(slides.slide_layouts[6]); hidden._element.set('show', '0')
    hidden.shapes.add_textbox(0, 0, Inches(2), Inches(1)).text = 'HIDDEN_SLIDE'
    end = slides.slides.add_slide(slides.slide_layouts[6]); end.shapes.add_textbox(0, 0, Inches(2), Inches(1)).text = 'ENDING'
    path = tmp_path/'slides.pptx'; slides.save(path)
    parsed = extract(path, 5000000, 256*1024**2)
    text = parsed['text']
    assert [text.index(c) for c in '①②③④'] == sorted(text.index(c) for c in '①②③④')
    assert text.index('④') < text.index('GROUP') < text.index('MERGED') < text.index('BOTTOM')
    assert text.count('MERGED') == 1 and 'NOTES_OMITTED' not in text and 'HIDDEN_SLIDE' not in text
    assert any('Slide 2: no readable text' in w for w in parsed['warnings'])
    assert segment(parsed)['sections'][-1]['page_start'] == 4


def test_office_rejections_and_limits(office_files, tmp_path):
    for path in office_files:
        with pytest.raises(ValueError, match='limit'): extract(path, 3, 256*1024**2)
        with pytest.raises(ValueError, match='limit'): extract(path, 5000000, 10)
    fake = tmp_path/'fake.docx'; fake.write_bytes(office_files[1].read_bytes())
    with pytest.raises(ValueError, match='extension'): validate_package(fake, 256*1024**2)
    fake.write_bytes(b'not a zip')
    with pytest.raises(ValueError, match='Invalid or encrypted'): validate_package(fake, 256*1024**2)
    fake.write_bytes(office_files[0].read_bytes())
    with zipfile.ZipFile(fake, 'a') as archive: archive.writestr('../escape', 'bad')
    with pytest.raises(ValueError, match='Unsafe'): validate_package(fake, 256*1024**2)
    fake.write_bytes(office_files[0].read_bytes())
    rewrite_member(fake, 'word/document.xml', lambda b: b.replace(b'<w:document', b'<!DOCTYPE x [<!ENTITY a "secret">]><w:document', 1))
    from defusedxml.common import DTDForbidden
    with pytest.raises(DTDForbidden): validate_package(fake, 256*1024**2)


def test_markdown_controls_pdf_heading_and_legacy_revision(tmp_path):
    path = tmp_path/'markers.md'
    path.write_text('# 标题\n正文\ue200cite\ue202turn1search2\ue201。普通引用[来源](https://example.com)\n', encoding="utf-8")
    parsed = extract(path, 5000000, 1000000)
    assert '\ue200' not in parsed['text'] and '来源' in parsed['text']
    headings = _plain_headings('1｜概览\n正文\n1｜概览\n续页\n1. 数字列表\n1|x| = 1\n2|y| = 2\n2｜结论\n', pdf=True)
    assert [h['title'] for h in headings] == ['1｜概览', '2｜结论']
    old = {'text': 'hello', 'text_hash': 'oldhash', 'version': 1}
    assert segment(old)['preview_revision'] == digest([1, 'oldhash', 'auto', 10000])
    old.pop('version')
    assert segment(old)['preview_revision'] == digest([1, 'oldhash', 'auto', 10000])


@pytest.mark.parametrize('index', [0, 1, 2])
def test_office_api_preview_synthesis_download(document_client, office_files, index):
    client, settings = document_client
    from audio_intel import db
    from audio_intel.worker import JobContext
    from tts.pipeline import process_job
    source = office_files[index]
    headers = {'Idempotency-Key': uuid.uuid4().hex}
    files = {'file': (source.name, source.read_bytes())}
    accepted = client.post('/api/v1/tts/document-imports', files=files, headers=headers)
    assert accepted.status_code == 202, accepted.text
    identifier = accepted.json()['id']
    assert client.post('/api/v1/tts/document-imports', files=files, headers=headers).status_code == 200
    deadline = time.monotonic() + 15
    while time.monotonic() < deadline:
        record = client.get('/api/v1/tts/document-imports/' + identifier).json()
        assert record['state'] != 'failed', record
        if record['state'] == 'ready': break
        time.sleep(.05)
    assert record['state'] == 'ready'
    preview = client.post(f'/api/v1/tts/document-imports/{identifier}/preview', json={}).json()
    assert preview['section_count'] == 2
    job, _ = submit_document(client, identifier, preview)
    result = process_job(JobContext(db.get_job(job['id']), 'test'))
    db.finish_job(job['id'], 'succeeded', result_json=result)
    response = client.get(f"/api/v1/jobs/{job['id']}/document/download?mode=sections")
    assert response.status_code == 200
    with zipfile.ZipFile(io.BytesIO(response.content)) as archive: assert len(archive.namelist()) == 2
    assert client.get(f"/api/v1/jobs/{job['id']}/document/download?mode=complete").status_code == 200
    caps = client.get('/api/v1/capabilities').json()['tts']['document_jobs']
    assert {'docx', 'xlsx', 'pptx'} <= set(caps['formats'])
    schema = client.get('/openapi.json').json()
    description = schema['paths']['/api/v1/tts/document-imports']['post']['description']
    assert all(fmt in description for fmt in ['DOCX', 'XLSX', 'PPTX'])


def test_excel_empty_cached_string_and_1904_dates(tmp_path):
    from openpyxl import Workbook
    from openpyxl.utils.datetime import CALENDAR_MAC_1904
    book = Workbook(); book.epoch = CALENDAR_MAC_1904
    sheet = book.active
    sheet.append(['=IF(1,"","")', dt.datetime(2026, 9, 11), dt.timedelta(hours=49, minutes=30)])
    sheet['B1'].number_format = 'yyyy-mm-dd'
    sheet['C1'].number_format = '[h]:mm:ss'
    path = tmp_path/'dates.xlsx'; book.save(path)
    rewrite_member(path, 'xl/worksheets/sheet1.xml', lambda b: b.replace(b'<c r="A1">', b'<c r="A1" t="str">'))
    parsed = extract(path, 5000000, 256*1024**2)
    assert '2026-09-11' in parsed['text'] and '49:30:00' in parsed['text']
    assert not parsed['warnings']


@pytest.mark.parametrize('index', [0, 1, 2])
def test_office_bearer_auth(tmp_path, monkeypatch, office_files, index):
    from dataclasses import replace
    from fastapi.testclient import TestClient
    from audio_intel import api, db
    from audio_intel.config import settings
    local = replace(settings, data_dir=tmp_path/'data', temp_dir=tmp_path/'tmp', api_key='office-test-key', mock_mode=True, min_free_disk_bytes=0)
    monkeypatch.setattr(api, 'settings', local)
    monkeypatch.setattr(db, 'settings', local)
    path = office_files[index]
    with TestClient(api.create_app()) as client:
        files = {'file': (path.name, path.read_bytes())}
        assert client.post('/api/v1/tts/document-imports', files=files, headers={'Idempotency-Key': uuid.uuid4().hex}).status_code == 401
        headers = {'Authorization': 'Bearer office-test-key', 'Idempotency-Key': uuid.uuid4().hex}
        accepted = client.post('/api/v1/tts/document-imports', files=files, headers=headers)
        assert accepted.status_code == 202
        assert client.post('/api/v1/tts/document-imports', files=files, headers=headers).status_code == 200
        assert client.get('/api/v1/tts/document-imports', headers=headers).status_code == 200


def test_office_async_failure_cleanup(document_client):
    client, local = document_client
    response = client.post('/api/v1/tts/document-imports', files={'file': ('broken.docx', b'encrypted or broken office')}, headers={'Idempotency-Key': uuid.uuid4().hex})
    assert response.status_code == 202
    identifier = response.json()['id']
    deadline = time.monotonic() + 15
    while time.monotonic() < deadline:
        record = client.get('/api/v1/tts/document-imports/' + identifier).json()
        if record['state'] == 'failed': break
        time.sleep(.05)
    assert record['state'] == 'failed'
    assert 'Invalid or encrypted' in record['error']
    assert not (local.data_dir/'documents'/identifier/'parsed.json').exists()
    assert client.delete('/api/v1/tts/document-imports/' + identifier).status_code == 204
    assert not (local.data_dir/'documents'/identifier).exists()
