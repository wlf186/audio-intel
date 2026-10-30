from __future__ import annotations

import asyncio
from copy import deepcopy
import io
import json
import os
import subprocess
import sys
import time
import httpx

from docx import Document
from fastapi.testclient import TestClient
from PIL import Image
from pypdf import PdfReader
import pytest

from audio_intel import api, db, optimization_api as routes, optimization_engine as engine, optimization_store as store, purge
from test_voiceprint_groups_exports import local


PROFILE = {'base_url': 'http://ollama.invalid:11434', 'use_proxy': False, 'model': 'test-model', 'enabled': True,
           'context_length': 32768, 'model_context_length': 131072, 'model_digest': 'test-digest', 'prompt_version': engine.PROMPT_VERSION, 'verified_at': '2026-10-23T00:00:00+00:00'}


def test_checkpoint_rejects_changed_fallback_and_does_not_double_count(configured):
    identifier = job(source(2))
    run, _ = store.create(identifier, 'checkpoint-review-key', False, PROFILE, engine.source_hash(source(2)), [], engine.prepare_units(source(2)))
    store.transition(run, 'running', ('queued',))
    with pytest.raises(ValueError, match='exact original'):
        store.checkpoint(run, [(0, 'unverified rewrite')], review_reason='review_failed')
    text = store.units(run)[0]['original_text']
    assert store.checkpoint(run, [(0, text)], review_reason='review_failed')
    assert store.checkpoint(run, [(0, text)], review_reason='review_failed')
    assert store.get(run)['needs_review_units'] == store.get(run)['completed_units'] == 1
    store.transition(run, 'paused', ('running',))
    assert not store.checkpoint(run, [(1, store.units(run)[1]['original_text'])], review_reason='review_failed')
    assert store.units(run)[1]['text'] is None


def test_storage_failure_never_becomes_completed_fallback(configured, monkeypatch):
    def fail(*args, **kwargs):
        raise OSError('disk unavailable')
    monkeypatch.setattr(store, 'checkpoint_part', fail)
    identifier = job(source(2))
    with TestClient(api.create_app()) as client:
        start(client, identifier)
        state = wait(client, identifier, lambda s: bool(s['candidate'] and s['candidate']['state']=='failed'))
        assert state['published'] is None
        assert state['candidate']['completed_units'] == state['candidate']['needs_review_units'] == 0


@pytest.mark.parametrize('code, expected, attempts', [(400, 'failed', 1), (429, 'paused', 3), (503, 'paused', 3)])
def test_provider_errors_do_not_become_paragraph_fallback(configured, monkeypatch, code, expected, attempts):
    calls = []
    async def generate(*args):
        calls.append(code)
        response = httpx.Response(code, request=httpx.Request('POST', 'http://model.invalid/api/chat'))
        response.raise_for_status()
    monkeypatch.setattr(engine, 'generate', generate)
    identifier = job(source(2))
    with TestClient(api.create_app()) as client:
        start(client, identifier)
        state = wait(client, identifier, lambda s: bool(s['candidate'] and s['candidate']['state']==expected))
        assert state['published'] is None
        assert state['candidate']['completed_units'] == state['candidate']['needs_review_units'] == 0
    assert len(calls) == attempts


def test_v15_migration_clears_legacy_optimization_preserving_source(configured):
    identifier = job(source(2))
    with TestClient(api.create_app()) as client:
        start(client, identifier)
        run = wait(client, identifier, lambda s: bool(s['published']))['published']['id']
    with db.connect() as connection:
        before = [tuple(row) for row in connection.execute('SELECT run_id,ordinal,source_json,optimized_text FROM optimization_units')]
        connection.execute('ALTER TABLE optimization_units DROP COLUMN review_status')
        connection.execute('ALTER TABLE optimization_units DROP COLUMN review_reason')
        connection.execute('ALTER TABLE transcript_optimizations DROP COLUMN needs_review_units')
        connection.execute('UPDATE schema_meta SET version=15')
    db.init_db()
    db.init_db()
    with db.connect() as connection:
        assert before and not connection.execute('SELECT 1 FROM optimization_units').fetchall()
        assert connection.execute('SELECT version FROM schema_meta').fetchone()[0] == 21
        assert not connection.execute('PRAGMA foreign_key_check').fetchall()
    assert store.status(identifier) == {'published': None, 'candidate': None}
    assert db.get_job(identifier)['result'] == source(2)


def docx_text(content):
    document = Document(io.BytesIO(content))
    return '\n'.join([p.text for p in document.paragraphs] + [cell.text for table in document.tables for row in table.rows for cell in row.cells])


def source(count=8):
    return {'language': 'Chinese', 'duration': count * 3, 'speakers': [
        {'id': 's1', 'name': 'Mike Smith', 'label': 'Mike Smith'}, {'id': 's2', 'name': 'Lucy Lee', 'label': 'Lucy Lee'}],
        'segments': [{'id': i, 'speaker': 's1' if i % 2 == 0 else 's2', 'speaker_label': 'Mike Smith' if i % 2 == 0 else 'Lucy Lee',
                      'start': i * 3, 'end': (i + 1) * 3, 'text': f'那个第{i}项，阿特拉斯预算15万元，不是50万元。暂时没有确认，周五先测试。'} for i in range(count)]}


def job(result=None):
    value = db.create_job('asr', '会议.wav', {'hotword_lists': [{'id': 'old', 'name': '历史热词', 'terms': ['Atlas']}]})
    db.update_job(value['id'], state='succeeded', result_json=result or source())
    return value['id']


def url(identifier):
    return f'/api/v1/jobs/{identifier}/transcript-optimization'


def wait(client, identifier, predicate, timeout=8):
    deadline = time.monotonic() + timeout
    while time.monotonic() < deadline:
        response = client.get(url(identifier))
        assert response.status_code == 200, response.text
        value = response.json()
        if predicate(value):
            return value
        time.sleep(.025)
    raise AssertionError(value)


@pytest.fixture
def configured(local, monkeypatch):
    store.save_settings(PROFILE)
    monkeypatch.setattr(purge, 'settings', local)
    async def check(self, profile):
        assert profile['model_digest'] == PROFILE['model_digest']
    async def generate(profile, batch, hotwords, before, after, activity):
        assert hotwords == ['Atlas']
        return [(p['ordinal'], p['original_text'].replace('那个', '').replace('阿特拉斯', 'Atlas')) for p in batch]
    monkeypatch.setattr(routes.Manager, 'check_model', check)
    monkeypatch.setattr(engine, 'generate', generate)
    async def ensure_context(profile, *, load=False):
        return profile['context_length']
    monkeypatch.setattr(engine, 'ensure_context', ensure_context)
    async def review(*args, **kwargs):
        return {'outcome': 'pass', 'details': {'stage': 'review', 'issues': []}}
    monkeypatch.setattr(engine, 'review', review)
    async def summary(*args, **kwargs):
        return ['暂未确认预算，周五先测试。']
    from audio_intel import optimization_summaries
    # These regressions exercise the saved v1 policy; v2 has its own fixtures.
    monkeypatch.setattr(optimization_summaries, 'VERSION', 1)
    monkeypatch.setattr(optimization_summaries, 'generate', summary)
    return local


def start(client, identifier, regenerate=False, key='optimization-test-key-0001'):
    return client.post(url(identifier), json={'regenerate': regenerate}, headers={'Idempotency-Key': key})


def test_configuration_verification_binding_and_auth(local, monkeypatch):
    async def verify(config):
        return {**PROFILE, **config}
    async def discover(config):
        return {'version': 'test', 'models': [{'name': 'test-model', 'digest': 'test-digest'}]}
    monkeypatch.setattr(engine, 'verify', verify)
    monkeypatch.setattr(engine, 'discover', discover)
    root = '/api/v1/settings/transcript-optimization'
    with TestClient(api.create_app()) as client:
        assert client.get(root).json()['enabled'] is False
        config = {key: PROFILE[key] for key in ('base_url', 'model', 'use_proxy', 'enabled')}
        assert client.put(root, json=config).status_code == 409
        assert client.post(root+'/discover', json={'base_url': 'file:///tmp/x'}).status_code == 422
        assert client.post(root+'/discover', json={'base_url': 'http://user:secret@localhost'}).status_code == 422
        assert client.post(root+'/discover', json={'base_url': config['base_url']}).json()['models'][0]['name'] == 'test-model'
        verified = client.post(root+'/verify', json={k: v for k, v in config.items() if k != 'enabled'}).json()
        config['verification_token'] = verified['verification_token']
        assert client.put(root, json={**config, 'use_proxy': True}).status_code == 409
        assert client.put(root, json=config).json()['enabled'] is True
        assert client.put(root, json={**config, 'enabled': False}).json()['enabled'] is False
        from dataclasses import replace
        with monkeypatch.context() as protected:
            protected.setattr(api, 'settings', replace(local, api_key='test-secret'))
            assert client.post(root+'/verify', json={'base_url': config['base_url'], 'model': 'x'}, headers={'Authorization':'Bearer invalid'}).status_code == 401
        schema = client.get('/openapi.json').json()
        assert root in schema['paths']
        assert 'text_version' in schema['components']['schemas']['TranscriptExportRequest']['properties']
        for model, fields in {'OptimizationSummary': ['needs_review_units'], 'OptimizedResult': ['needs_review_units'], 'OptimizedParagraph': ['review_status', 'review_reason', 'candidate_text', 'review_details'], 'ReviewDetails': ['stage', 'issues'], 'ReviewIssue': ['source_quote', 'candidate_quote', 'reason']}.items():
            assert all(field in schema['components']['schemas'][model]['properties'] for field in fields)


def test_units_source_coverage_scene_splits_and_long_speech():
    result = source(2)
    result['segments'] = [{'id': 0, 'speaker':'s1', 'speaker_label':'Mike Smith', 'start':0, 'end':6, 'text':'第一段。第二段。',
                           'words':[{'text':'第一段','start':0,'end':2},{'text':'第二段','start':3,'end':5}]}]
    result['video'] = {'scenes':[{'id':'a','start':0,'end':3},{'id':'b','start':3,'end':6}]}
    before = deepcopy(result)
    units = engine.prepare_units(result)
    assert [p['original_text'] for p in units] == ['第一段。', '第二段。']
    assert [(p['sources'][0]['start_offset'],p['sources'][0]['end_offset']) for p in units] == [(0,4), (4,8)]
    assert result == before
    result.pop('video')
    result['segments'][0].update(text='长句内容。'*900, words=[])
    units = engine.prepare_units(result)
    assert ''.join(p['original_text'] for p in units) == result['segments'][0]['text']
    assert len(units) > 1 and all(len(unit['original_text']) <= 400 for unit in units)
    chunks = engine.prepare_parts(PROFILE, units[0], [])
    assert all(len(part['original_text']) <= 1200 for part in chunks)
    assert ''.join(part['original_text'] for part in chunks) == units[0]['original_text']
    assert ''.join(s['text'] for p in units for s in p['sources']) == result['segments'][0]['text']
    with pytest.raises(ValueError, match='窗口'):
        engine.validate_budget({**PROFILE,'context_length':16384}, units, ['热词'*10000])


def test_v14_migration_preserves_history_queue_and_idempotency(local):
    history = job()
    queued, replayed = db.create_job_idempotent('tts', 'historical queue', {'text': '保存快照'}, None, 'tts', 'existing-key', 'existing-request')
    assert not replayed
    before = {identifier: deepcopy(db.get_job(identifier)) for identifier in [history, queued['id']]}
    with db.connect() as connection:
        receipt = tuple(connection.execute('SELECT * FROM job_idempotency').fetchone())
        sequence = tuple(connection.execute('SELECT * FROM queue_sequence').fetchone())
        for table in ['optimization_parts', 'optimization_units', 'transcript_optimizations', 'optimization_settings']:
            connection.execute('DROP TABLE ' + table)
        connection.execute('UPDATE schema_meta SET version=14')
    db.init_db()
    db.init_db()
    assert {identifier: db.get_job(identifier) for identifier in before} == before
    assert not store.settings()['enabled']
    with db.connect() as connection:
        assert connection.execute('SELECT min(version) FROM schema_meta').fetchone()[0] == 21
        assert tuple(connection.execute('SELECT * FROM job_idempotency').fetchone()) == receipt
        assert tuple(connection.execute('SELECT * FROM queue_sequence').fetchone()) == sequence
        assert not connection.execute('PRAGMA foreign_key_check').fetchall()
    repeated, replayed = db.create_job_idempotent('tts', 'historical queue', {'text': '保存快照'}, None, 'tts', 'existing-key', 'existing-request')
    assert replayed and repeated['id'] == queued['id']


def test_accepted_key_replays_after_disabling_or_changing_configuration(configured):
    identifier = job()
    with TestClient(api.create_app()) as client:
        assert start(client, identifier).status_code == 202
        run = wait(client, identifier, lambda state: bool(state['published']))['published']['id']
        store.save_settings({'base_url': 'http://other.invalid', 'model': 'other', 'use_proxy': False, 'enabled': False})
        replayed = start(client, identifier)
        assert replayed.status_code == 200
        assert replayed.json()['published']['id'] == run
        conflict = start(client, identifier, True)
        assert conflict.status_code == 409 and 'Idempotency' in conflict.json()['detail']
        assert start(client, identifier, True, 'new-key-after-disable').status_code == 409


def test_complete_idempotency_summary_exports_latest_and_cleanup(configured):
    identifier = job()
    before = deepcopy(db.get_job(identifier))
    with TestClient(api.create_app()) as client:
        response = start(client, identifier)
        assert response.status_code == 202, response.text
        assert start(client, identifier).status_code == 200
        assert start(client, identifier, True).status_code == 409
        state = wait(client, identifier, lambda s: bool(s['published']))
        first = state['published']['id']
        assert state['published']['progress'] == 1
        result = client.get(url(identifier)+'/result', params={'run_id':first}).json()
        assert result['complete'] and len(result['paragraphs']) == 8
        assert 'Atlas' in result['paragraphs'][0]['text']
        assert '阿特拉斯' in result['paragraphs'][0]['original_text']
        response = client.get(url(identifier))
        assert client.get(url(identifier),headers={'If-None-Match':response.headers['etag']}).status_code == 304
        summaries = client.get('/api/v1/jobs').json()
        assert 'optimization' in json.dumps(summaries)
        assert 'original_text' not in json.dumps(summaries) and 'hotwords' not in json.dumps(summaries)
        export = f'/api/v1/jobs/{identifier}/transcript-export'
        body = {'format':'docx','text_version':'comparison','optimization_id':first,'redact':True,'replacements':[{'source':'Atlas','target':'平台'},{'source':'阿特拉斯','target':'平台'}]}
        response = client.post(export,json=body)
        assert response.status_code == 200, response.text
        document = Document(io.BytesIO(response.content))
        text = docx_text(response.content)
        assert '原文 / Original' in text and '平台' in text and 'Mike Smith' not in text
        assert len(document.tables) == 8
        assert document.tables[0].cell(0,0).text == '原文 / Original'
        assert document.sections[0].page_width > document.sections[0].page_height
        assert document.styles['Normal'].font.size.pt == 11
        assert 'Atlas' in result['paragraphs'][0]['text']
        response = client.post(export,json={**body,'format':'pdf'})
        assert response.status_code == 200
        pdf = '\n'.join(p.extract_text() for p in PdfReader(io.BytesIO(response.content)).pages)
        assert '原文' in pdf and '平台' in pdf and 'Mike Smith' not in pdf
        assert client.post(export,json={**body,'format':'srt'}).status_code == 422
        assert client.post(export,json={'format':'docx','text_version':'comparison'}).status_code == 422
        assert start(client,identifier,True,'optimization-second-key').status_code == 202
        state = wait(client,identifier,lambda s: bool(s['published'] and s['published']['id']!=first))
        assert client.get(url(identifier)+'/result',params={'run_id':first}).status_code == 404
        assert client.delete(url(identifier)).status_code == 409
        assert client.delete(url(identifier)+'?purge=true').status_code == 204
        assert client.get(url(identifier)).json()['published'] is None
        with db.connect() as connection:
            assert connection.execute('SELECT COUNT(*) FROM optimization_units').fetchone()[0] == 0
        assert db.get_job(identifier) == before


def test_pause_resume_disable_and_no_late_writes(configured, monkeypatch):
    calls = []
    async def generate(profile,batch,hotwords,before,after,activity):
        calls.append([p['ordinal'] for p in batch])
        if len(calls) == 2:
            await asyncio.sleep(60)
        return [(p['ordinal'],p['original_text'].replace('那个','')) for p in batch]
    monkeypatch.setattr(engine,'generate',generate)
    identifier = job(source(100))
    with TestClient(api.create_app()) as client:
        start(client,identifier)
        state = wait(client,identifier,lambda s: bool(s['candidate'] and s['candidate']['completed_units']>0))
        run = state['candidate']['id']
        assert client.post(url(identifier)+'/pause',json={'run_id':run}).status_code == 200
        paused = client.get(url(identifier)).json()['candidate']
        assert paused['state'] == 'paused' and 0 < paused['progress'] < 1
        assert not store.checkpoint(run,[(99,'late output')])
        assert client.post(f'/api/v1/jobs/{identifier}/transcript-export',json={'format':'pdf','text_version':'optimized','optimization_id':run}).status_code == 409
        assert client.post(url(identifier)+'/resume',json={'run_id':run}).status_code == 200
        state = wait(client,identifier,lambda s: bool(s['published']))
        assert state['published']['completed_units'] == 100
        assert not any(set(calls[0])&set(batch) for batch in calls[1:])
        start(client,identifier,True,'optimization-regenerate-key')
        disabled = {k: PROFILE[k] for k in ('base_url','model','use_proxy')}
        assert client.put('/api/v1/settings/transcript-optimization',json={**disabled,'enabled':False}).status_code == 200
        state = client.get(url(identifier)).json()
        if state['candidate']:
            assert state['candidate']['state']=='paused'
            assert client.post(url(identifier)+'/resume',json={'run_id':state['candidate']['id']}).status_code == 409
        assert state['published']
        assert client.delete(url(identifier)+'?purge=true').status_code == 204
        assert not store.checkpoint(run,[(0,'resurrected')])


def test_failed_regeneration_retains_published_and_parent_purge(configured,monkeypatch):
    identifier = job()
    with TestClient(api.create_app()) as client:
        start(client,identifier)
        first = wait(client,identifier,lambda s:bool(s['published']))['published']['id']
        async def fail(*args):
            raise ValueError('Model failed')
        monkeypatch.setattr(engine,'generate',fail)
        start(client,identifier,True,'optimization-failed-generation')
        state=wait(client,identifier,lambda s:bool(s['candidate'] and s['candidate']['state']=='failed'))
        assert state['published']['id']==first
        response=client.delete(f'/api/v1/jobs/{identifier}?purge=true')
        assert response.status_code==204,response.text
        with db.connect() as connection:
            assert connection.execute('SELECT COUNT(*) FROM transcript_optimizations').fetchone()[0]==0
            assert connection.execute('SELECT COUNT(*) FROM optimization_units').fetchone()[0]==0


def test_single_paragraph_checkpoints_and_restart_recovery(configured,monkeypatch):
    calls=[]
    async def small(profile,batch,*args):
        calls.append(len(batch))
        if len(batch)>2:
            raise engine.InvalidOutput('length')
        return [(p['ordinal'],p['original_text']) for p in batch]
    monkeypatch.setattr(engine,'generate',small)
    identifier=job()
    with TestClient(api.create_app()) as client:
        start(client,identifier)
        state=wait(client,identifier,lambda s:bool(s['published']))
        assert state['published']['completed_units']==8 and max(calls)==1
        run=state['published']['id']
    with db.connect() as connection:
        connection.execute("UPDATE transcript_optimizations SET published=0,state='running' WHERE id=?",(run,))
    with TestClient(api.create_app()) as client:
        state=wait(client,identifier,lambda s:bool(s['candidate'] and s['candidate']['state']=='paused'))
        assert state['candidate']['completed_units']==8


def test_video_comparison_images_redaction_and_scene_reassignment(configured):
    result=source(4)
    result['video']={'revision':1,'duration':12,'time_origin':0,'scenes':[
        {'id':'a','start':0,'end':6,'frame_time':1,'image':'a.jpg'},
        {'id':'b','start':6,'end':12,'frame_time':8,'image':'b.jpg'}],'export_hidden_scene_ids':['b']}
    identifier=job(result)
    frames=configured.jobs_dir/identifier/'output'/'frames';frames.mkdir(parents=True)
    for name in ('a.jpg','b.jpg'):
        Image.new('RGB',(80,60),'blue').save(frames/name)
    with TestClient(api.create_app()) as client:
        start(client,identifier)
        run=wait(client,identifier,lambda s:bool(s['published']))['published']['id']
        response=client.post(f'/api/v1/jobs/{identifier}/transcript-export',json={'format':'docx','text_version':'comparison','optimization_id':run})
        assert response.status_code==200,response.status_code
        document=Document(io.BytesIO(response.content))
        assert len(document.inline_shapes)==1
        assert sum(len(table.columns)==2 for table in document.tables)==4
        result['video']['scenes'][0]['end']=9;result['video']['scenes'][1]['start']=9
        db.update_job(identifier,result_json=result)
        assert client.get(url(identifier)+'/result',params={'run_id':run}).status_code == 409


def test_queue_limit_and_settings_are_independent(local):
    identifiers=[job() for _ in range(6)]
    for identifier in identifiers[:5]:
        store.create(identifier,'key-'+identifier,False,PROFILE,'hash',[],engine.prepare_units(source()))
    from fastapi import HTTPException
    with pytest.raises(HTTPException) as error:
        store.create(identifiers[5],'key-six',False,PROFILE,'hash',[],engine.prepare_units(source()))
    assert error.value.status_code==429
    assert error.value.headers.get('Retry-After')


def test_omitted_units_are_completed_and_exported_with_explicit_comparison(configured,monkeypatch):
    async def omit(profile,batch,*args):return [(p['ordinal'],'') for p in batch]
    monkeypatch.setattr(engine,'generate',omit)
    data=source(2)
    for segment in data['segments']:segment['text']='呃。'
    identifier=job(data)
    with TestClient(api.create_app()) as client:
        start(client,identifier)
        state=wait(client,identifier,lambda s:bool(s['published']))
        run=state['published']['id']
        assert state['published']['progress']==1 and state['published']['completed_units']==2
        result=client.get(url(identifier)+'/result',params={'run_id':run}).json()
        assert all(p['omitted'] and p['text']=='' for p in result['paragraphs'])
        endpoint=f'/api/v1/jobs/{identifier}/transcript-export'
        text=client.post(endpoint,json={'format':'txt','text_version':'optimized','optimization_id':run,'redact':False})
        assert text.status_code==200 and '呃' not in text.text
        response=client.post(endpoint,json={'format':'docx','text_version':'comparison','optimization_id':run,'redact':False})
        assert response.status_code==200
        content=docx_text(response.content)
        assert '已省略纯占位语' in content and '呃。' in content


def test_old_candidate_regeneration_retains_published_and_connection_profile(configured):
    identifier=job()
    with TestClient(api.create_app()) as client:
        start(client,identifier)
        published=wait(client,identifier,lambda s:bool(s['published']))['published']['id']
        legacy={**PROFILE,'prompt_version':1}
        # Simulate a candidate checkpoint created before the policy upgrade.
        candidate,_=store.create(identifier,'old-candidate',True,legacy,engine.source_hash(source()),['Atlas'],engine.prepare_units(source()))
        store.transition(candidate,'paused',('queued',))
        store.save_settings(legacy)
        state=client.get(url(identifier)).json()
        assert state['candidate']['requires_regeneration']
        response=start(client,identifier,True,'new-policy-key')
        assert response.status_code==202,response.text
        state=wait(client,identifier,lambda s:bool(s['published'] and s['published']['id']!=published))
        assert store.get(state['published']['id'])['config']['prompt_version']==engine.PROMPT_VERSION
        assert store.settings()['prompt_version']==1


def test_cli_regeneration_waits_for_new_draft_and_preserves_omission_contract(monkeypatch,capsys):
    import runpy
    import urllib.request
    old={'id':'old','state':'succeeded'}
    pending={'published':old,'candidate':{'id':'new','state':'running'}}
    responses=[pending,pending,{'published':{'id':'new','state':'succeeded'},'candidate':None},
               {'run_id':'new','paragraphs':[{'text':'','omitted':True,'original_text':'呃。'}]}]
    calls=[]
    def urlopen(request,timeout):
        calls.append(request)
        return io.BytesIO(json.dumps(responses.pop(0)).encode())
    monkeypatch.setattr(urllib.request,'urlopen',urlopen)
    monkeypatch.setattr(time,'sleep',lambda _:None)
    monkeypatch.setattr(sys,'argv',['optimize_transcript.py','--url','http://test','--job','job','--regenerate'])
    runpy.run_path('scripts/optimize_transcript.py',run_name='__main__')
    assert json.loads(calls[0].data)=={'regenerate':True}
    assert calls[-1].full_url.endswith('run_id=new')
    assert json.loads(capsys.readouterr().out)['paragraphs'][0]['omitted']


@pytest.mark.parametrize('mode', ['polish', 'rewrite', 'editorial'])
def test_rewrite_modes_snapshot_and_idempotency(configured, mode):
    identifier = job(source(2))
    with TestClient(api.create_app()) as client:
        response = client.post(url(identifier), json={'rewrite_mode': mode}, headers={'Idempotency-Key':'mode-snapshot-test'})
        assert response.status_code == 202
        published = wait(client, identifier, lambda s: bool(s['published']))['published']
        assert published['rewrite_mode'] == mode
        assert store.get(published['id'])['config']['rewrite_mode'] == mode
        result = client.get(url(identifier)+'/result', params={'run_id': published['id']}).json()
        assert result['rewrite_mode'] == mode
        assert client.post(url(identifier), json={'rewrite_mode':mode}, headers={'Idempotency-Key':'mode-snapshot-test'}).status_code == 200
        other = 'polish' if mode != 'polish' else 'editorial'
        assert client.post(url(identifier), json={'rewrite_mode':other}, headers={'Idempotency-Key':'mode-snapshot-test'}).status_code == 409
        assert client.post(url(identifier), json={'rewrite_mode':'unknown'}, headers={'Idempotency-Key':'invalid-mode-key'}).status_code == 422
        caps = client.get('/api/v1/capabilities').json()['asr']['transcript_optimization']
        assert caps['rewrite_modes'] == ['polish','rewrite','editorial']
        assert caps['default_rewrite_mode'] == 'rewrite'


def test_legacy_idempotency_replay_without_mode_still_works(configured):
    identifier = job(source(2))
    legacy = {**PROFILE, 'prompt_version':3}
    legacy.pop('rewrite_mode', None)
    run, _ = store.create(identifier, 'legacy-mode-key', False, legacy, engine.source_hash(source(2)), [], engine.prepare_units(source(2)))
    store.transition(run, 'paused', ('queued',))
    with TestClient(api.create_app()) as client:
        assert client.post(url(identifier), json={}, headers={'Idempotency-Key':'legacy-mode-key'}).status_code == 200
        assert client.post(url(identifier), json={'rewrite_mode':'rewrite'}, headers={'Idempotency-Key':'legacy-mode-key'}).status_code == 409
        assert store.get(run)['config'] == legacy


@pytest.mark.parametrize('failures', [0, 1, 2, 3])
def test_verification_format_retries_are_bounded(monkeypatch, failures):
    calls = []
    async def remote(connection, path, payload=None):
        if path == 'api/tags':
            return {'models': [{'name': 'test-model', 'digest': 'test-digest'}]}
        if path == 'api/show':
            return {'capabilities': ['completion'], 'model_info': {'test.context_length': 32768}}
        return {'models': [{'name': 'test-model', 'digest': 'test-digest', 'context_length': 32768}]}
    async def generate(*args):
        calls.append(1)
        if len(calls) <= failures:
            raise engine.InvalidOutput('invalid')
        return [(0, '预算15万元。')]
    async def sleep(delay): pass
    monkeypatch.setattr(engine, 'remote_json', remote)
    monkeypatch.setattr(engine, 'generate', generate)
    monkeypatch.setattr(engine.asyncio, 'sleep', sleep)
    if failures == 3:
        with pytest.raises(engine.InvalidOutput, match='已自动重试两次'):
            asyncio.run(engine.verify(PROFILE))
    else:
        assert asyncio.run(engine.verify(PROFILE))['verified_at']
    assert len(calls) == min(failures + 1, 3)


def test_v16_migration_clears_legacy_candidate_preserving_source(configured):
    identifier=job(source(1))
    run,_=store.create(identifier,'migration-candidate-key',False,PROFILE,'hash',[],engine.prepare_units(source(1)))
    before=store.units(run)
    with db.connect() as connection:
        connection.execute('ALTER TABLE optimization_units DROP COLUMN candidate_text')
        connection.execute('ALTER TABLE optimization_units DROP COLUMN review_details_json')
        connection.execute('UPDATE schema_meta SET version=16')
    db.init_db()
    assert before and store.units(run) == []
    assert db.get_job(identifier)['result'] == source(1)
    with db.connect() as connection:
        assert connection.execute('SELECT version FROM schema_meta').fetchone()[0]==21
