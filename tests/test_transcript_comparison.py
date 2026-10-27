from pathlib import Path
import io

from docx import Document
from pypdf import PdfReader

from audio_intel.transcript_pdf import render as pdf
from audio_intel.transcript_docx import render as docx
from audio_intel.transcript_diff import text_changes


def test_comparison_landscape_columns_long_rows_and_literal_markup(tmp_path):
    original = '原文段落开头。' + '预算15万元，不是50万元，审批后才执行。' * 90 + '原文段落结尾<&>。'
    edited = '改写段落开头。' + '预算为15万元，并非50万元，执行前须审批。' * 35 + '改写段落结尾<&>。'
    result = {'_optimized': True, 'text_version': 'comparison', 'segments': [
        {'speaker': 's1', 'speaker_label': '甲', 'start': 0, 'end': 20, 'text': edited, 'original_text': original, 'review_status': 'needs_review'},
        {'speaker': 's2', 'speaker_label': '乙', 'start': 20, 'end': 21, 'text': '最后一段', 'original_text': '最后一段'},
    ]}
    data = pdf('对照', result, tmp_path)
    reader = PdfReader(io.BytesIO(data))
    assert len(reader.pages) >= 2
    assert all(p.mediabox.width > p.mediabox.height for p in reader.pages)
    text = '\n'.join(p.extract_text() for p in reader.pages)
    for marker in ('原文段落开头', '原文段落结尾', '改写段落开头', '改写段落结尾', '最后一段', 'Needs review'):
        assert marker in text
    # Both columns are identified on continuation pages.
    assert sum('Original' in p.extract_text() for p in reader.pages) >= 2
    document = Document(io.BytesIO(docx('对照', result, tmp_path)))
    assert document.sections[0].page_width > document.sections[0].page_height
    assert len(document.tables) == 2
    table = document.tables[0]
    assert table.cell(1, 0).text == original and table.cell(1, 1).text == edited
    assert 'tblHeader' in table.rows[0]._tr.xml
    assert 'cantSplit' not in table.rows[1]._tr.xml
    assert any(run.font.strike for run in table.cell(1, 0).paragraphs[0].runs)
    assert table.columns[0].width == table.columns[1].width


def test_literal_diff_reconstructs_both_sides_and_bounds_large_inputs():
    for left, right in [('原文<&>\n预算15', '改写<&>\n预算15'), ('呃', ''), ('同文', '同文'), ('字'*2000, '改'*2000)]:
        a, b = text_changes(left, right)
        assert ''.join(text for text, _ in a) == left
        assert ''.join(text for text, _ in b) == right
