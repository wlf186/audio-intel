"""Bounded, offline OOXML extraction into the document spoken-text contract."""
from __future__ import annotations

import datetime as dt
import posixpath
import re
import zipfile
from decimal import Decimal, ROUND_HALF_UP, localcontext
from pathlib import Path
from typing import Any

from defusedxml import ElementTree as ET

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
S = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
R = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'


def validate_package(path: Path, archive_limit: int) -> None:
    """Preflight before third-party parsers materialize XML or shared strings."""
    expected = {'.docx': 'word/document.xml', '.xlsx': 'xl/workbook.xml', '.pptx': 'ppt/presentation.xml'}
    try:
        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            names = {e.filename for e in entries}
            if len(entries) > 10000 or sum(e.file_size for e in entries) > archive_limit:
                raise ValueError('Office expanded size or member count exceeds the limit / Office 文件展开大小或成员数超限')
            if len(names) != len(entries):
                raise ValueError('Office archive contains duplicate members')
            for entry in entries:
                name = entry.filename
                if name.startswith('/') or '\\' in name or '..' in name.split('/') or ':' in name:
                    raise ValueError('Unsafe Office member path')
                if entry.flag_bits & 1:
                    raise ValueError('Encrypted Office documents are unsupported / 请先解除文档加密')
                if name.endswith(('.xml', '.rels')):
                    # Reject DTD/entity payloads even when the downstream library
                    # uses a different XML implementation. Clear nodes as we go.
                    with archive.open(entry) as source:
                        for _, node in ET.iterparse(source, forbid_dtd=True):
                            node.clear()
            if expected[path.suffix.lower()] not in names or '[Content_Types].xml' not in names:
                raise ValueError('Office content does not match the file extension / 文件内容与扩展名不符')
            types = ET.fromstring(archive.read('[Content_Types].xml'))
            main = next((e.get('ContentType', '') for e in types if e.get('PartName') == '/' + expected[path.suffix.lower()]), '')
            wanted = {'.docx': 'wordprocessingml.document.main+xml', '.xlsx': 'spreadsheetml.sheet.main+xml', '.pptx': 'presentationml.presentation.main+xml'}
            if not main.endswith(wanted[path.suffix.lower()]):
                raise ValueError('Unsupported Office package type / 请另存为 DOCX、XLSX 或 PPTX')
    except zipfile.BadZipFile as exc:
        raise ValueError('Invalid or encrypted Office document / Office 文件损坏或已加密') from exc


class SpokenText:
    def __init__(self, limit: int):
        self.limit = limit
        self.parts: list[str] = []
        self.length = 0
        self.headings: list[dict[str, Any]] = []
        self.warnings: list[str] = []

    def append(self, text: str, *, title: str = '', level: int = 1, basis: str = 'heading') -> None:
        text = text.strip()
        if not text:
            return
        if self.length + len(text) + 2 > self.limit:
            raise ValueError('Document text exceeds the configured limit')
        if title:
            self.headings.append({'offset': self.length, 'title': title[:200], 'level': level, 'basis': basis})
        self.parts.append(text + '\n\n')
        self.length += len(text) + 2

    def result(self, title: str) -> dict[str, Any]:
        return {'text': ''.join(self.parts), 'headings': self.headings, 'warnings': list(dict.fromkeys(self.warnings)), 'title': title}


def docx_text(path: Path, limit: int) -> dict[str, Any]:
    from docx import Document
    from docx.text.paragraph import Paragraph
    document = Document(path)
    output = SpokenText(limit)
    fields: list[bool] = []
    has_toc = any(re.match(r'\s*TOC\b', n.text or '') for n in document.element.iter(W + 'instrText')) or any(re.match(r'\s*TOC\b', n.get(W + 'instr', '')) for n in document.element.iter(W + 'fldSimple'))

    def visible(node: Any) -> str:
        tag = node.tag
        if tag in {W + 'del', W + 'moveFrom'}:
            return ''
        if tag == W + 'r' and node.find(W + 'rPr/' + W + 'vanish') is not None:
            return ''
        if tag in {W + 'drawing', W + 'pict', W + 'object'}:
            output.warnings.append('Word: images/embedded objects omitted / 图片或嵌入对象未朗读')
            return ''
        if tag == W + 'fldChar':
            kind = node.get(W + 'fldCharType')
            if kind == 'begin':
                fields.append(False)
            elif kind == 'end' and fields:
                fields.pop()
            return ''
        if tag == W + 'instrText':
            if fields and re.match(r'\s*TOC\b', node.text or ''):
                fields[-1] = True
            return ''
        if tag == W + 'fldSimple' and re.match(r'\s*TOC\b', node.get(W + 'instr', '')):
            return ''
        if tag == W + 't':
            return '' if any(fields) else node.text or ''
        if tag in {W + 'tab', W + 'br', W + 'cr'}:
            return '' if any(fields) else ('\t' if tag == W + 'tab' else '\n')
        return ''.join(visible(child) for child in node)

    def paragraph(node: Any) -> tuple[str, int | None]:
        text = visible(node)
        p = Paragraph(node, document)
        style = p.style
        if has_toc and ((style is not None and re.match(r'^TOC\s*\d*$', style.name or '', re.I)) or text.strip() in {'目录', 'Contents', 'Table of Contents'}):
            return '', None
        properties = node.find(W + 'pPr')
        outline = properties.find(W + 'outlineLvl') if properties is not None else None
        seen: set[str] = set()
        while outline is None and style is not None and style.style_id not in seen:
            seen.add(style.style_id)
            properties = style.element.find(W + 'pPr')
            outline = properties.find(W + 'outlineLvl') if properties is not None else None
            style = style.base_style
        level = int(outline.get(W + 'val', '9')) + 1 if outline is not None else None
        return text, level if level is not None and level <= 9 else None

    def table(node: Any) -> str:
        rows = []
        for row in node.findall(W + 'tr'):
            values = []
            for cell in row.findall(W + 'tc'):
                merge = cell.find(W + 'tcPr/' + W + 'vMerge')
                if merge is not None and merge.get(W + 'val') != 'restart':
                    continue
                parts = []
                for child in cell:
                    if child.tag == W + 'p':
                        parts.append(paragraph(child)[0])
                    elif child.tag == W + 'tbl':
                        parts.append(table(child))
                values.append(' '.join(t.strip() for t in parts if t.strip()))
            if any(values):
                rows.append('；'.join(values))
        return '\n'.join(rows)

    def blocks(parent: Any) -> None:
        for node in parent:
            if node.tag == W + 'p':
                text, level = paragraph(node)
                if level and (len(text.strip()) > 100 or text.rstrip().endswith('。')):
                    output.warnings.append('Word: prose styled as a heading retained as body text / 过长或句式标题按正文保留')
                    level = None
                output.append(text, title=text if level else '', level=level or 1)
            elif node.tag == W + 'tbl':
                output.append(table(node))
            elif node.tag in {W + 'sdt', W + 'sdtContent', W + 'ins', W + 'moveTo'}:
                blocks(node)
    blocks(document.element.body)
    if has_toc:
        output.warnings.append('Word: automatic table of contents omitted / 自动目录未重复朗读')
    return output.result(document.core_properties.title or path.stem)


def display_cell(cell: Any, warnings: list[str], location: str) -> str:
    """Render common Excel display formats without evaluating any formulas."""
    value = cell.value
    if value is None:
        return ''
    if isinstance(value, bool):
        return 'TRUE' if value else 'FALSE'
    if cell.data_type == 'e':
        warnings.append(f'{location}: Excel error {value} / 单元格错误')
        return f'单元格错误 {value}'
    fmt = cell.number_format or 'General'
    if isinstance(value, dt.datetime):
        return value.isoformat(sep=' ', timespec='seconds') if re.search(r'[hs]', fmt, re.I) else value.date().isoformat()
    if isinstance(value, (dt.date, dt.time)):
        return value.isoformat()
    if isinstance(value, dt.timedelta):
        total = int(value.total_seconds())
        return f'{total // 3600:02}:{abs(total) % 3600 // 60:02}:{abs(total) % 60:02}'
    if not isinstance(value, (int, float)):
        return str(value)
    if fmt.lower() == 'general' or fmt == '@':
        return str(value) if not isinstance(value, float) else format(value, '.15g')
    sections = fmt.split(';')
    selected = sections[1] if value < 0 and len(sections) > 1 else sections[2] if value == 0 and len(sections) > 2 else sections[0]
    selected = re.sub(r'\[\$([^\]-]*)-[^\]]+\]', r'\1', selected)
    selected = re.sub(r'\[(?:Black|Blue|Cyan|Green|Magenta|Red|White|Yellow|Color\d+)\]', '', selected, flags=re.I)
    selected = re.sub(r'_.|\*.', '', selected)
    match = re.fullmatch(r'([^0#?\[\]]*)([#,0]+)(?:\.([0#]+))?(%?)([^0#?\[\]]*)', selected)
    if match:
        prefix, integer, decimals, percent, suffix = match.groups()
        # Quoted and escaped literals belong to the displayed value, not the mask.
        literal = lambda s: re.sub(r'\\(.)', r'\1', s.replace('"', ''))
        digits = len(decimals or '')
        with localcontext() as context:
            context.prec = 340
            number = Decimal(str(abs(value) if value < 0 and len(sections) > 1 else value))
            if percent:
                number *= 100
            scaling = len(integer) - len(integer.rstrip(','))
            if scaling:
                number /= 1000 ** scaling
                integer = integer.rstrip(',')
            number = number.quantize(Decimal(1).scaleb(-digits), rounding=ROUND_HALF_UP)
            rendered = format(number, f',.{digits}f' if ',' in integer else f'.{digits}f')
        if ',' not in integer:
            whole, dot, fraction = rendered.partition('.')
            sign = '-' if whole.startswith('-') else ''
            rendered = sign + whole.lstrip('-').zfill(integer.count('0')) + dot + fraction
        if decimals and decimals.endswith('#'):
            minimum = len(decimals.rstrip('#'))
            whole, dot, fraction = rendered.partition('.')
            fraction = fraction.rstrip('0')
            fraction = fraction.ljust(minimum, '0')
            rendered = whole + (dot + fraction if fraction else '')
        return literal(prefix) + rendered + percent + literal(suffix)
    warnings.append(f'{location}: unsupported number format {fmt}; raw value used / 数字格式不支持，显示原始值')
    return str(value)


def xlsx_text(path: Path, limit: int) -> dict[str, Any]:
    from openpyxl import load_workbook
    from openpyxl.cell.read_only import ReadOnlyCell
    from openpyxl.worksheet._reader import WorkSheetParser
    from openpyxl.utils.cell import get_column_letter, range_boundaries

    class SparseParser(WorkSheetParser):
        # The pinned parser yields only physical cells. Public iter_rows pads
        # holes up to the last row/column, which is unsafe for sparse uploads.
        def parse_cell(self, element: Any) -> dict[str, Any]:
            result = super().parse_cell(element)
            cached = element.find(S + 'v')
            empty_string = element.get('t') == 'str' and cached is not None
            result['missing_formula'] = element.find(S + 'f') is not None and result['value'] is None and not empty_string
            return result

    output = SpokenText(limit)
    workbook = load_workbook(path, read_only=True, data_only=True, keep_links=False)
    try:
        for sheet in workbook.worksheets:
            if sheet.sheet_state != 'visible':
                output.warnings.append(f'{sheet.title}: hidden worksheet omitted / 隐藏工作表未朗读')
                continue
            start = output.length
            headers: dict[int, str] = {}
            previous = 0
            formal: list[tuple[int, int, int, int, dict[int, str]]] = []
            # Table definitions are not loaded by openpyxl's read-only mode.
            member = sheet._worksheet_path
            merged = []
            with sheet._get_source() as source:
                for _, element in ET.iterparse(source):
                    if element.tag == S + 'mergeCell':
                        merged.append(range_boundaries(element.get('ref')))
                    element.clear()
            rels = posixpath.join(posixpath.dirname(member), '_rels', posixpath.basename(member) + '.rels')
            if rels in workbook._archive.namelist():
                for rel in ET.fromstring(workbook._archive.read(rels)):
                    if rel.get('Type', '').endswith('/vmlDrawing'):
                        output.warnings.append(f'{sheet.title}: annotations/legacy drawing omitted / 批注或旧式图形未朗读')
                    elif rel.get('Type', '').endswith('/drawing'):
                        output.warnings.append(f'{sheet.title}: images/charts omitted / 图片或图表未朗读')
                    if rel.get('TargetMode') == 'External' or not rel.get('Type', '').endswith('/table'):
                        continue
                    target = rel.get('Target', '')
                    name = target.lstrip('/') if target.startswith('/') else posixpath.normpath(posixpath.join(posixpath.dirname(member), target))
                    if name.startswith('../') or '\\' in name:
                        raise ValueError('Unsafe Office relationship')
                    table = ET.fromstring(workbook._archive.read(name))
                    if table.get('headerRowCount', '1') == '0':
                        continue
                    left, top, right, bottom = range_boundaries(table.get('ref'))
                    columns = table.find(S + 'tableColumns')
                    formal.append((left, top, right, bottom, {left + i: c.get('name', '') for i, c in enumerate(columns if columns is not None else [])}))
            with sheet._get_source() as source:
                parser = SparseParser(source, sheet._shared_strings, data_only=True, epoch=workbook.epoch, date_formats=workbook._date_formats, timedelta_formats=workbook._timedelta_formats)
                for number, raw in parser.parse():
                    if parser.row_dimensions.get(str(number), {}).get('hidden') in {'1', 'true', True}:
                        output.warnings.append(f'{sheet.title}: hidden rows omitted / 隐藏行未朗读')
                        continue
                    if number != previous + 1:
                        headers = {}
                    previous = number
                    values: dict[int, str] = {}
                    cells = []
                    for data in raw:
                        missing = data.pop('missing_formula')
                        col = data['column']
                        if any(left <= col <= right and top <= number <= bottom and (col, number) != (left, top) for left, top, right, bottom in merged):
                            continue
                        if any(d.get('hidden') in {'1', 'true', True} and int(d['min']) <= col <= int(d['max']) for d in parser.column_dimensions.values()):
                            output.warnings.append(f'{sheet.title}: hidden columns omitted / 隐藏列未朗读')
                            continue
                        cell = ReadOnlyCell(sheet, **data)
                        location = f'{sheet.title}!{cell.coordinate}'
                        if missing:
                            value = '公式结果缺失'
                            output.warnings.append(f'{location}: formula result missing; recalculate and save / 公式结果缺失，请重新计算并保存')
                        else:
                            value = display_cell(cell, output.warnings, location).strip()
                        if value:
                            values[col] = value
                            cells.append(cell)
                    if not values:
                        headers = {}
                        continue
                    matching = [t for t in formal if t[1] <= number <= t[3]]
                    declared_header = any(number == t[1] for t in matching)
                    styled_header = len(cells) >= 2 and all(isinstance(c.value, str) and len(c.value) <= 100 and c.font.bold for c in cells) and len(set(values.values())) == len(values)
                    if declared_header or styled_header:
                        headers = dict(values)
                        for left, top, right, bottom in merged:
                            if top == number and left in headers:
                                headers.update({col: headers[left] for col in range(left, right + 1)})
                        for left, top, right, bottom, labels in matching:
                            headers.update(labels)
                        # Retain the header itself, including header-only tables.
                        output.append('；'.join(values.values()))
                        continue
                    row_headers = dict(headers)
                    for left, top, right, bottom, labels in matching:
                        row_headers.update(labels)
                    if row_headers:
                        text = '；'.join(f'{row_headers.get(c, "列 " + get_column_letter(c))}：{v}' for c, v in values.items())
                    else:
                        ordered = list(values.values())
                        text = ordered[0] + ('：' + '；'.join(ordered[1:]) if len(ordered) > 1 else '')
                    output.append(text)
            if output.length > start:
                output.headings.append({'offset': start, 'title': sheet.title[:200], 'level': 1, 'basis': 'worksheet'})
            else:
                output.warnings.append(f'{sheet.title}: no readable cells / 无可读单元格')
    finally:
        workbook.close()
    return output.result(path.stem)


def pptx_text(path: Path, limit: int) -> dict[str, Any]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER
    presentation = Presentation(path)
    output = SpokenText(limit)
    offsets = []
    for number, slide in enumerate(presentation.slides, 1):
        offsets.append(output.length)
        if slide._element.get('show') in {'0', 'false'}:
            output.warnings.append(f'Slide {number}: hidden slide omitted / 隐藏幻灯片未朗读')
            continue
        units: list[tuple[float, float, float, int, str]] = []

        def collect(shapes: Any, sx: float = 1, sy: float = 1, dx: float = 0, dy: float = 0) -> None:
            for shape in shapes:
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    transform = shape._element.grpSpPr.xfrm
                    if transform is None or not transform.chExt.cx or not transform.chExt.cy:
                        output.warnings.append(f'Slide {number}: invalid group geometry / 组合形状坐标无效')
                        continue
                    gx, gy = shape.width / transform.chExt.cx, shape.height / transform.chExt.cy
                    collect(shape.shapes, sx * gx, sy * gy, dx + sx * (shape.left - gx * transform.chOff.x), dy + sy * (shape.top - gy * transform.chOff.y))
                    continue
                if shape.is_placeholder and shape.placeholder_format.type in {PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.SLIDE_NUMBER}:
                    continue
                if shape.has_table:
                    rows = []
                    for row in shape.table.rows:
                        values = [cell.text.replace('\v', '\n').strip() for cell in row.cells if not cell.is_spanned]
                        if any(values):
                            rows.append('；'.join(values))
                    text = '\n'.join(rows)
                elif shape.has_text_frame:
                    text = shape.text.replace('\v', '\n').strip()
                else:
                    output.warnings.append(f'Slide {number}: image/chart/embedded object omitted / 图片、图表或嵌入对象未朗读')
                    continue
                if text:
                    units.append((dy + sy * shape.top, dx + sx * shape.left, sy * shape.height, len(units), text))
        collect(slide.shapes)
        units.sort(key=lambda u: (u[0], u[1], u[3]))
        ordered = []
        index = 0
        while index < len(units):
            first = units[index]
            index += 1
            band = [first]
            while index < len(units) and abs(units[index][0] - first[0]) <= min(91440, .25 * min(first[2], units[index][2])):
                band.append(units[index])
                index += 1
            ordered.extend(sorted(band, key=lambda u: (u[1], u[3])))
        if not ordered:
            output.warnings.append(f'Slide {number}: no readable text; OCR may be required / 此页无可读文字，可能需要 OCR')
            continue
        title_shape = slide.shapes.title
        title = title_shape.text.strip() if title_shape is not None else ordered[0][4].splitlines()[0]
        start = output.length
        for unit in ordered:
            output.append(unit[4])
        output.headings.append({'offset': start, 'title': f'{number}. {title}'[:200], 'level': 1, 'basis': 'slide'})
    result = output.result(presentation.core_properties.title or path.stem)
    result['page_offsets'] = offsets
    return result


def extract_office(path: Path, max_chars: int, archive_limit: int) -> dict[str, Any]:
    validate_package(path, archive_limit)
    return {'.docx': docx_text, '.xlsx': xlsx_text, '.pptx': pptx_text}[path.suffix.lower()](path, max_chars)
