"""Offline, editable concise transcripts with bounded in-memory generation."""
from __future__ import annotations

import hashlib
import io
import json
import threading
import time
from collections import OrderedDict
from pathlib import Path
from typing import Any

from .transcript import reading_sections
from .transcript_diff import text_changes
from .utils import timecode

_lock = threading.Lock()
_render_lock = threading.Lock()
_cache: OrderedDict[str, tuple[float, bytes]] = OrderedDict()
_cache_generation = 0
_cache_limit = 64 * 1024**2
_max_image_bytes = 128 * 1024**2
_max_text_chars = 4_000_000


def render(title: str, result: dict[str, Any], output_dir: Path) -> bytes:
    from docx import Document
    from docx.image.image import Image
    from docx.enum.section import WD_ORIENT
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor

    if sum(len(segment.get('text', '')) + len(segment.get('original_text', '')) for segment in result.get('segments', [])) > _max_text_chars:
        raise ValueError('Transcript exceeds DOCX text budget; split the source')
    document = Document()
    properties = document.core_properties
    for name in ('author', 'last_modified_by', 'subject', 'keywords', 'comments', 'category', 'identifier'):
        setattr(properties, name, '')
    properties.title = title
    section = document.sections[0]
    comparison = result.get('text_version') == 'comparison'
    section.orientation = WD_ORIENT.LANDSCAPE if comparison else WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = (Mm(297), Mm(210)) if comparison else (Mm(210), Mm(297))
    section.left_margin = section.right_margin = Mm(18)
    section.top_margin = section.bottom_margin = Mm(18)
    for name in ('Normal', 'Title', 'Heading 1'):
        style = document.styles[name]
        style.font.name = 'Calibri'
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '宋体')
    normal = document.styles['Normal']
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.2
    document.add_heading(title, 0)
    hidden = set((result.get('video') or {}).get('export_hidden_scene_ids') or [])
    sections = reading_sections(result)
    image_bytes = 0
    frame_root = (output_dir / 'frames').resolve()
    for index, item in enumerate(sections):
        if result.get('video'):
            document.add_heading(f"画面 {index + 1} · {timecode(item['start'])} – {timecode(item['end'])}", 1)
            if not item.get('image') or item['id'] in hidden:
                cell = document.add_table(rows=1, cols=1).cell(0, 0)
                cell.text = '无稳定画面 / No stable picture' if not item.get('image') else '画面已隐藏 / Picture hidden'
                shade = OxmlElement('w:shd')
                shade.set(qn('w:fill'), 'F0F2F4')
                cell._tc.get_or_add_tcPr().append(shade)
            else:
                name = item['image']
                path = (frame_root / name).resolve()
                if Path(name).name != name or not path.is_relative_to(frame_root):
                    raise ValueError('Invalid frame path')
                image_bytes += path.stat().st_size
                if image_bytes > _max_image_bytes:
                    raise ValueError('Images exceed DOCX resource budget; split the source')
                # Read the validated bytes, avoiding filename metadata in the package.
                stream = io.BytesIO(path.read_bytes())
                picture = Image.from_file(stream)
                scale = min(int(section.page_width - section.left_margin - section.right_margin) / picture.width, int(Mm(100)) / picture.height)
                stream.seek(0)
                document.add_picture(stream, width=int(picture.width * scale), height=int(picture.height * scale))
        for paragraph in item['paragraphs']:
            label = document.add_paragraph()
            label.paragraph_format.keep_with_next = True
            label.add_run(f"{paragraph.get('speaker_label') or paragraph['speaker']}　{timecode(paragraph['start'])} – {timecode(paragraph['end'])}").bold = True
            if paragraph.get('review_status') == 'needs_review':
                label.add_run('　保留原文·待核对 / Original retained · Needs review').italic = True
            if 'original_text' in paragraph:
                table = document.add_table(rows=2, cols=2)
                table.autofit = False
                column_width = int((section.page_width - section.left_margin - section.right_margin) / 2)
                for column in table.columns:
                    column.width = column_width
                for row in table.rows:
                    for cell in row.cells:
                        cell.width = column_width
                        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
                table.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
                for cell, title in zip(table.rows[0].cells, ('原文 / Original', '改写 / Rewritten')):
                    cell.paragraphs[0].add_run(title).bold = True
                    shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'F0F2F4')
                    cell._tc.get_or_add_tcPr().append(shade)
                changes = text_changes(paragraph['original_text'], paragraph['text'])
                for index, (cell, parts) in enumerate(zip(table.rows[1].cells, changes)):
                    content = cell.paragraphs[0]
                    content.paragraph_format.keep_together = False
                    content.paragraph_format.keep_with_next = False
                    for text, changed in parts:
                        run = content.add_run(text)
                        if changed:
                            run.font.color.rgb = RGBColor.from_string('9C3A32' if index == 0 else '12634B')
                            run.font.strike = index == 0
                document.add_paragraph().paragraph_format.space_after = Pt(4)
            else:
                document.add_paragraph(paragraph['text'])
        if not item['paragraphs']:
            document.add_paragraph(('此段无优化正文 / No optimized text in this section' if result.get('_optimized') else '此段无识别语音 / No recognized speech'))
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def download(title: str, result: dict[str, Any], output_dir: Path, cache_variant: str = "") -> bytes:
    key = hashlib.sha256(json.dumps([str(output_dir), title, result, cache_variant], ensure_ascii=False, sort_keys=True).encode()).hexdigest()
    now = time.monotonic()
    with _lock:
        generation = _cache_generation
        for old in [k for k, (expiry, _) in _cache.items() if expiry <= now]:
            del _cache[old]
        cached = _cache.get(key)
        if cached:
            _cache.move_to_end(key)
            return cached[1]
    if not _render_lock.acquire(blocking=False):
        raise BlockingIOError('DOCX renderer is busy; retry shortly')
    try:
        payload = render(title, result, output_dir)
        with _lock:
            if generation == _cache_generation and len(payload) <= _cache_limit:
                while _cache and sum(len(data) for _, data in _cache.values()) + len(payload) > _cache_limit:
                    _cache.popitem(last=False)
                _cache[key] = (time.monotonic() + 300, payload)
        return payload
    finally:
        _render_lock.release()
